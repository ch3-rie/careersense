from datetime import datetime, timedelta, timezone
from io import BytesIO

from docx import Document
from openpyxl import load_workbook

from app.models import Account, AlumniAward, AlumniProfile, FurtherStudy, TracerSubmission, UniversityRecord
from app.security import hash_password
from app.services.reports import build_oaaps_report, classify_further_study, latest_submissions, pct
from tests.conftest import auth_header, login

COLLEGE = "College of OAAPS Report"
PROGRAM = "Bachelor of Science in Report Testing"
YEAR = "2025"


def test_reports_count_latest_per_alumni(client, db_session):
    maria = db_session.query(Account).filter(Account.personal_email == "maria.reyes@gmail.com").one()
    existing = db_session.query(TracerSubmission).filter(TracerSubmission.account_id == maria.id).count()
    db_session.add(
        TracerSubmission(
            account_id=maria.id,
            student_id=maria.linked_student_id,
            data_json={
                "is_currently_employed": "Yes",
                "present_related_degree": "Yes",
                "pres_occ": "Software Engineer",
            },
            extra_answers={},
            alignment_status="Aligned",
            alignment_score=100,
        )
    )
    db_session.commit()
    latest = latest_submissions(db_session)
    maria_rows = [row for row in latest if row.account_id == maria.id]
    assert len(maria_rows) == 1
    assert existing >= 1

    admin = login(client, "admin@auf.edu.ph", "Admin@AUF2026")
    reports = client.get("/api/admin/reports", headers=auth_header(admin["access_token"]))
    assert reports.status_code == 200
    body = reports.json()
    assert body["alumni_with_submissions"] == len(latest)
    assert body["total_submissions"] >= body["alumni_with_submissions"]
    assert body["currently_employed_submissions"] <= body["alumni_with_submissions"]


def test_pct_guards_zero_denominator():
    assert pct(5, 0) == 0.0
    assert pct(1, 2) == 50.0
    assert pct(1, 3) == 33.33


def test_classify_further_study_labels():
    assert classify_further_study("Master in Business Administration", False) == "pursuing_masters"
    assert classify_further_study("PhD in Education", True) == "completed_doctorate"
    assert classify_further_study("Postgraduate Diploma in Data Science", False) == "postgrad_cert"
    assert classify_further_study("Bachelor of Laws", False) == "second_degree"
    assert classify_further_study("Short course", False) == "other"


def _record(db, student_id, email, first, last, year=YEAR):
    row = UniversityRecord(
        student_id=student_id,
        first_name=first,
        last_name=last,
        personal_email=email,
        degree=PROGRAM,
        year_graduated=year,
        course_code="BSRT",
        college=COLLEGE,
    )
    db.add(row)
    db.flush()
    return row


def _alumni(db, record, email, payload, extra_studies=None, submitted_at=None, alignment="Aligned"):
    account = Account(
        personal_email=email,
        password_hash=hash_password("Alumni@2026"),
        linked_student_id=record.student_id,
        role="Alumni",
        status="Active",
        is_verified=True,
        privacy_consent=True,
    )
    db.add(account)
    db.flush()
    db.add(
        AlumniProfile(
            account_id=account.id,
            first_name=record.first_name,
            last_name=record.last_name,
            degree=record.degree,
            year_graduated=record.year_graduated,
        )
    )
    db.add(
        TracerSubmission(
            account_id=account.id,
            student_id=record.student_id,
            data_json=payload,
            extra_answers={},
            alignment_status=alignment,
            alignment_score=100 if alignment == "Aligned" else 0,
            submitted_at=submitted_at,
        )
    )
    for item in extra_studies or []:
        db.add(FurtherStudy(account_id=account.id, **item))
    db.flush()
    return account


def _oaaps_fixture(db_session):
    employed_studies = _record(db_session, "OAAPS-0001", "oaaps.one@example.com", "One", "Productive")
    _alumni(
        db_session,
        employed_studies,
        "oaaps.one@example.com",
        {
            "first_name": "One",
            "last_name": "Productive",
            "degree": PROGRAM,
            "year_graduated": YEAR,
            "is_currently_employed": "Yes",
            "present_job_is_first": "Yes",
            "first_related": "Yes",
            "first_stat": "Regular/Permanent",
            "time_to_first_job": "Less than a month",
            "enroll_further_studies": "Yes",
            "further_studies": [
                {
                    "course_degree": "Master of Science in Data Analytics",
                    "school": "AUF",
                    "year_enrolled": "2026",
                    "scholarship": "CHED Scholar",
                    "is_graduated": "No",
                },
                {
                    "course_degree": "Certificate in Project Management",
                    "school": "AUF",
                    "year_enrolled": "2025",
                    "scholarship": "",
                    "is_graduated": "No",
                },
            ],
        },
        extra_studies=[
            {
                "course_degree": "Master of Science in Data Analytics",
                "school": "AUF",
                "year_enrolled": "2026",
                "scholarship": "CHED Scholar",
                "is_graduated": False,
            },
            {
                "course_degree": "Certificate in Project Management",
                "school": "AUF",
                "year_enrolled": "2025",
                "scholarship": "",
                "is_graduated": False,
            },
        ],
    )
    seeking = _record(db_session, "OAAPS-0002", "oaaps.two@example.com", "Two", "Seeking")
    _alumni(
        db_session,
        seeking,
        "oaaps.two@example.com",
        {
            "first_name": "Two",
            "last_name": "Seeking",
            "degree": PROGRAM,
            "year_graduated": YEAR,
            "is_currently_employed": "No",
            "reason_current": ["Looking for a job, but cannot find one"],
            "enroll_further_studies": "No",
        },
        alignment="Unknown",
    )
    self_employed = _record(db_session, "OAAPS-0003", "oaaps.three@example.com", "Three", "Owner")
    _alumni(
        db_session,
        self_employed,
        "oaaps.three@example.com",
        {
            "first_name": "Three",
            "last_name": "Owner",
            "degree": PROGRAM,
            "year_graduated": YEAR,
            "is_currently_employed": "Yes",
            "present_job_is_first": "Yes",
            "first_stat": "Self-Employed",
            "first_related": "No",
            "time_to_first_job": "7 to 11 months",
            "enroll_further_studies": "No",
        },
        alignment="Misaligned",
    )
    _record(db_session, "OAAPS-0004", "oaaps.four@example.com", "Four", "Untraced")
    db_session.commit()


def test_oaaps_report_formulas_and_no_double_count(client, db_session):
    _oaaps_fixture(db_session)
    admin = login(client, "admin@auf.edu.ph", "Admin@AUF2026")
    response = client.get(
        "/api/admin/reports/oaaps",
        params={"college": COLLEGE, "program": PROGRAM, "year": YEAR},
        headers=auth_header(admin["access_token"]),
    )
    assert response.status_code == 200
    body = response.json()
    assert body["title"] == "OVERALL GRADUATE PRODUCTIVITY REPORT"
    assert body["header"]["total_graduates"] == 4
    assert body["header"]["graduates_traced"] == 3
    assert body["profile"]["productive"] == 2
    assert body["profile"]["not_yet_engaged"] == 1
    assert body["overall_productivity_rate"] == 66.67
    assert body["workforce"]["key_rate"] == 66.67
    assert body["professional"]["key_rate"] == 33.33
    assert body["academic"]["key_rate"] == 33.33
    assert body["distribution"]["workforce"] == 66.67

    by_indicator = {row["indicator"]: row for row in body["executive_summary"]}
    assert by_indicator["Graduates with Multiple Productivity Indicators"]["count"] == 2
    assert by_indicator["Employed Graduates"]["count"] == 2
    assert by_indicator["Graduates Pursuing Further Studies"]["count"] == 1
    assert by_indicator["Graduates with Professional Advancement"]["count"] == 1

    status = {row["indicator"]: row for row in body["workforce"]["employment_status"]}
    assert status["Employed"]["count"] == 1
    assert status["Self-Employed / Entrepreneur"]["count"] == 1
    assert status["Freelance / Project-Based"]["available"] is False
    assert status["Unemployed – Seeking Employment"]["count"] == 1
    assert status["Total Traced"]["count"] == 3

    relevance = {row["indicator"]: row for row in body["workforce"]["employment_relevance"]}
    assert relevance["Employed in a field related to degree"]["count"] == 1
    assert relevance["Employed in a different field"]["count"] == 1
    assert relevance["Employed locally"]["available"] is False

    quality = {row["indicator"]: row for row in body["workforce"]["employment_quality"]}
    assert quality["Permanent/Regular"]["count"] == 1
    assert quality["Full-Time Employment"]["available"] is False

    time_rows = {row["indicator"]: row for row in body["workforce"]["time_to_employment"]}
    assert time_rows["Employed before graduation"]["available"] is False
    assert time_rows["Employed within 3 months"]["count"] == 1
    assert time_rows["Employed within 12 months"]["count"] == 1

    education = {row["indicator"]: row for row in body["academic"]["further_education"]}
    assert education["Pursuing Master’s Degree"]["count"] == 1
    assert education["Pursuing Postgraduate Certificate/Diploma"]["count"] == 1

    achievements = {row["indicator"]: row for row in body["academic"]["achievements"]}
    assert achievements["Received Scholarship/Fellowship"]["count"] == 1
    assert achievements["Published Research"]["available"] is False

    assert body["professional"]["achievements_available"] is False
    assert all(row["target"] is None and row["status"] is None for row in body["scorecard"])
    assert body["by_program"][-1]["college_program"] == "TOTAL"
    assert body["by_program"][-1]["graduates"] == 4
    assert body["by_program"][-1]["traced"] == 3


def test_oaaps_latest_submission_and_period_filter(client, db_session):
    college = "College of OAAPS Latest Filter"
    record = UniversityRecord(
        student_id="OAAPS-0100",
        first_name="Late",
        last_name="Update",
        personal_email="oaaps.latest@example.com",
        degree=PROGRAM,
        year_graduated="2024",
        course_code="BSRT",
        college=college,
    )
    db_session.add(record)
    db_session.flush()
    account = _alumni(
        db_session,
        record,
        "oaaps.latest@example.com",
        {
            "is_currently_employed": "Yes",
            "present_job_is_first": "Yes",
            "first_stat": "Regular/Permanent",
            "enroll_further_studies": "No",
        },
        submitted_at=datetime.now(timezone.utc) - timedelta(days=10),
    )
    db_session.add(
        TracerSubmission(
            account_id=account.id,
            student_id=record.student_id,
            data_json={
                "is_currently_employed": "No",
                "reason_current": ["Continuing Education"],
                "enroll_further_studies": "No",
            },
            extra_answers={},
            alignment_status="Unknown",
            alignment_score=50,
        )
    )
    db_session.commit()

    report = build_oaaps_report(db_session, college=college, program=PROGRAM, year="2024")
    assert report["header"]["graduates_traced"] == 1
    status = {row["indicator"]: row for row in report["workforce"]["employment_status"]}
    assert status["Employed"]["count"] == 0
    assert status["Not Seeking Employment"]["count"] == 1

    admin = login(client, "admin@auf.edu.ph", "Admin@AUF2026")
    empty_period = client.get(
        "/api/admin/reports/oaaps",
        params={
            "college": college,
            "year": "2024",
            "period_from": "2010-01-01",
            "period_to": "2010-12-31",
        },
        headers=auth_header(admin["access_token"]),
    )
    assert empty_period.status_code == 200
    empty = empty_period.json()
    assert empty["header"]["graduates_traced"] == 0
    assert empty["overall_productivity_rate"] == 0.0
    assert empty["header"]["total_graduates"] == 1


def test_oaaps_empty_cohort_and_invalid_dates(client, db_session):
    admin = login(client, "admin@auf.edu.ph", "Admin@AUF2026")
    headers = auth_header(admin["access_token"])
    empty = client.get("/api/admin/reports/oaaps", params={"year": "1999"}, headers=headers)
    assert empty.status_code == 200
    body = empty.json()
    assert body["header"]["total_graduates"] == 0
    assert body["header"]["graduates_traced"] == 0
    assert body["overall_productivity_rate"] == 0.0
    assert body["workforce"]["employment_status"][-1]["percentage"] == 0.0

    bad = client.get("/api/admin/reports/oaaps", params={"period_from": "13-13-2024"}, headers=headers)
    assert bad.status_code == 400
    inverted = client.get(
        "/api/admin/reports/oaaps",
        params={"period_from": "2026-12-31", "period_to": "2026-01-01"},
        headers=headers,
    )
    assert inverted.status_code == 400


def _report_labels(body):
    labels = ["Total Graduates", "Traced Graduates", "Productive Graduates", "Not Yet Productively Engaged"]
    labels.extend(row["indicator"] for row in body["executive_summary"])
    for key in ("employment_status", "employment_relevance", "employment_quality", "time_to_employment"):
        labels.extend(row["indicator"] for row in body["workforce"][key])
    labels.extend(row["indicator"] for row in body["professional"]["career_progression"])
    labels.extend(row["indicator"] for row in body["professional"]["credentials"])
    labels.extend(row["indicator"] for row in body["academic"]["further_education"])
    labels.extend(row["indicator"] for row in body["academic"]["achievements"])
    labels.extend(row["key_indicator"] for row in body["scorecard"])
    labels.extend(["Workforce Integration", "Professional Advancement", "Academic Advancement"])
    return labels


def test_directory_lists_match_productivity_counts(client, db_session):
    if db_session.get(UniversityRecord, "OAAPS-0001") is None:
        _oaaps_fixture(db_session)
    admin = login(client, "admin@auf.edu.ph", "Admin@AUF2026")
    headers = auth_header(admin["access_token"])
    params = {"college": COLLEGE, "program": PROGRAM, "year": YEAR}
    report = client.get("/api/admin/reports/oaaps", params=params, headers=headers)
    assert report.status_code == 200
    body = report.json()
    counts = {}
    for section in (
        body["executive_summary"],
        body["workforce"]["employment_status"],
        body["workforce"]["employment_relevance"],
        body["workforce"]["employment_quality"],
        body["workforce"]["time_to_employment"],
        body["professional"]["career_progression"],
        body["professional"]["credentials"],
        body["academic"]["further_education"],
        body["academic"]["achievements"],
    ):
        for row in section:
            counts[row["indicator"]] = None if row.get("available") is False else row["count"]

    def listing(indicator, **extra):
        response = client.get(
            "/api/admin/directory",
            params={**params, "cohort": "productivity", "indicator": indicator, "page_size": 50, **extra},
            headers=headers,
        )
        assert response.status_code == 200, response.text
        return response.json()

    for label in _report_labels(body):
        payload = listing(label)
        assert payload["title"] == label
        expected = counts.get(label)
        if expected is None and label not in {
            "Traced Graduates",
            "Productive Graduates",
            "Not Yet Productively Engaged",
            "Workforce Integration",
            "Professional Advancement",
            "Academic Advancement",
        } and label not in {row["key_indicator"] for row in body["scorecard"]}:
            if label in counts and counts[label] is None:
                assert payload["snapshot_total"] == 0
            elif label in counts:
                assert payload["snapshot_total"] == counts[label]
        elif label in counts and counts[label] is not None:
            assert payload["snapshot_total"] == counts[label]

    assert listing("Total Graduates")["snapshot_total"] == 4
    assert listing("Graduates Traced")["snapshot_total"] == 3
    assert listing("Not Yet Productively Engaged")["snapshot_total"] == 1
    assert listing("Freelance / Project-Based")["snapshot_total"] == 0
    assert "Four Untraced" in {row["name"] for row in listing("Total Graduates")["items"]}
    assert "Four Untraced" not in {row["name"] for row in listing("Graduates Traced")["items"]}

    future = client.get(
        "/api/admin/reports/oaaps",
        params={**params, "period_from": "2099-01-01"},
        headers=headers,
    )
    assert future.status_code == 200
    assert future.json()["header"]["total_graduates"] == 4
    assert future.json()["header"]["graduates_traced"] == 0
    assert listing("Total Graduates", period_from="2099-01-01")["snapshot_total"] == 4
    assert listing("Graduates Traced", period_from="2099-01-01")["snapshot_total"] == 0

    unknown = client.get(
        "/api/admin/directory",
        params={**params, "cohort": "productivity", "indicator": "Not a real category"},
        headers=headers,
    )
    assert unknown.status_code == 400


def test_oaaps_word_export_and_csv_still_work(client, db_session):
    admin = login(client, "admin@auf.edu.ph", "Admin@AUF2026")
    headers = auth_header(admin["access_token"])
    word = client.get("/api/admin/reports/oaaps/export", headers=headers)
    assert word.status_code == 200
    assert "wordprocessingml" in word.headers["content-type"]
    assert word.content[:2] == b"PK"
    csv_export = client.get("/api/admin/reports/export", headers=headers)
    assert csv_export.status_code == 200
    assert csv_export.headers["content-type"].startswith("text/csv")
    assert csv_export.text.splitlines()[0].startswith("email")


def test_report_filters_lists_and_exports_stay_consistent(client, db_session):
    if db_session.get(UniversityRecord, "OAAPS-0001") is None:
        _oaaps_fixture(db_session)
    awarded = db_session.query(Account).filter(Account.personal_email == "oaaps.one@example.com").one()
    awarded.profile.city = "Angeles City"
    db_session.add(AlumniAward(account_id=awarded.id, title="Dean's Lister", issuer="AUF"))
    if db_session.get(UniversityRecord, "OAAPS-0005") is None:
        unknown_record = _record(db_session, "OAAPS-0005", "oaaps.five@example.com", "Five", "Unstated")
        _alumni(
            db_session,
            unknown_record,
            "oaaps.five@example.com",
            {"first_name": "Five", "last_name": "Unstated", "degree": PROGRAM, "year_graduated": YEAR},
        )
    if not any(award.title == "Dean's Lister" for award in awarded.awards):
        db_session.add(AlumniAward(account_id=awarded.id, title="Dean's Lister", issuer="AUF"))
    db_session.commit()

    admin = login(client, "admin@auf.edu.ph", "Admin@AUF2026")
    headers = auth_header(admin["access_token"])
    scope = {"college": COLLEGE, "program": PROGRAM, "year": YEAR}

    employed = client.get("/api/admin/reports/employment", params={**scope, "employment": "employed"}, headers=headers)
    assert employed.status_code == 200, employed.text
    assert employed.json()["employed"] == employed.json()["total"] == 2
    assert employed.json()["unemployed"] == 0
    assert employed.json()["unknown"] == 0

    unemployed = client.get("/api/admin/reports/employment", params={**scope, "employment": "unemployed"}, headers=headers)
    assert unemployed.json()["unemployed"] == 1
    assert unemployed.json()["employed"] == 0
    assert unemployed.json()["total"] == 1

    missing = client.get("/api/admin/reports/employment", params={**scope, "employment": "retired"}, headers=headers)
    assert missing.status_code == 400

    located = client.get("/api/admin/reports/employment", params={**scope, "location": "Angeles City"}, headers=headers)
    assert located.json()["total"] == 1
    assert located.json()["employed"] == 1
    blank_location = client.get("/api/admin/reports/employment", params={**scope, "location": "__none__"}, headers=headers)
    assert blank_location.json()["total"] == 3
    assert "oaaps.one@example.com" not in blank_location.text

    csv_export = client.get("/api/admin/reports/export", params={**scope, "employment": "unemployed"}, headers=headers)
    assert csv_export.status_code == 200
    csv_lines = [line for line in csv_export.text.splitlines() if line.strip()]
    assert csv_lines[0].startswith("email")
    assert len(csv_lines) - 1 == unemployed.json()["total"]
    assert "oaaps.two@example.com" in csv_export.text
    assert "oaaps.one@example.com" not in csv_export.text
    assert "oaaps.five@example.com" not in csv_export.text

    recorded_awards = client.get(
        "/api/admin/reports/oaaps",
        params={**scope, "award": "__recorded__"},
        headers=headers,
    )
    assert recorded_awards.status_code == 200, recorded_awards.text
    assert recorded_awards.json()["header"]["total_graduates"] == 1
    assert recorded_awards.json()["header"]["graduates_traced"] == 1
    assert recorded_awards.json()["header"]["awards"] == "Has recorded awards"

    no_awards = client.get("/api/admin/reports/oaaps", params={**scope, "award": "__none__"}, headers=headers)
    assert no_awards.json()["header"]["total_graduates"] == 4
    assert no_awards.json()["header"]["graduates_traced"] == 3

    studies = client.get(
        "/api/admin/reports/oaaps",
        params={**scope, "further": "Master of Science in Data Analytics"},
        headers=headers,
    )
    assert studies.json()["header"]["total_graduates"] == 1
    assert studies.json()["header"]["graduates_traced"] == 1
    none_studies = client.get("/api/admin/reports/oaaps", params={**scope, "further": "__none__"}, headers=headers)
    assert none_studies.json()["header"]["total_graduates"] == 4
    assert none_studies.json()["profile"]["traced"] == 3
    combined = client.get(
        "/api/admin/reports/oaaps",
        params={**scope, "award": "__recorded__", "further": "__none__"},
        headers=headers,
    )
    assert combined.json()["header"]["total_graduates"] == 0
    assert combined.json()["header"]["graduates_traced"] == 0

    listing = client.get(
        "/api/admin/directory",
        params={**scope, "cohort": "productivity", "indicator": "Total Graduates", "award": "__recorded__", "page_size": 50},
        headers=headers,
    )
    assert listing.status_code == 200, listing.text
    assert listing.json()["snapshot_total"] == recorded_awards.json()["header"]["total_graduates"]
    assert listing.json()["items"][0]["name"] == "One Productive"

    employed_list = client.get(
        "/api/admin/directory",
        params={**scope, "cohort": "analytics", "answer": "employed", "employment": "employed", "page_size": 50},
        headers=headers,
    )
    assert employed_list.json()["snapshot_total"] == employed.json()["employed"]

    csv_list = client.get(
        "/api/admin/directory/export",
        params={**scope, "cohort": "productivity", "indicator": "Total Graduates", "award": "__recorded__", "format": "csv"},
        headers=headers,
    )
    assert csv_list.status_code == 200
    assert csv_list.text.splitlines()[0].startswith("Name,Student ID")
    assert len(csv_list.text.splitlines()) - 1 == listing.json()["snapshot_total"]

    workbook = client.get(
        "/api/admin/directory/export",
        params={**scope, "cohort": "analytics", "answer": "unemployed", "employment": "unemployed", "format": "xlsx"},
        headers=headers,
    )
    assert workbook.status_code == 200
    assert "spreadsheetml" in workbook.headers["content-type"]
    sheet = load_workbook(BytesIO(workbook.content)).active
    assert [cell.value for cell in sheet[1]][:3] == ["Name", "Student ID", "Email"]
    assert sheet.max_row - 1 == unemployed.json()["total"]

    word = client.get("/api/admin/reports/oaaps/export", params={**scope, "award": "__recorded__"}, headers=headers)
    cells = [cell.text for table in Document(BytesIO(word.content)).tables for row in table.rows for cell in row.cells]
    assert "Awards Received" in cells
    assert "Has recorded awards" in cells
    assert "1" in cells

    alumni = login(client, "liza.torres@gmail.com", "Alumni@2026")
    denied = client.get("/api/admin/directory/export", headers=auth_header(alumni["access_token"]))
    assert denied.status_code == 403
    assert client.get("/api/admin/directory/export").status_code == 401


def test_alumni_cannot_access_oaaps_report(client):
    alumni = login(client, "liza.torres@gmail.com", "Alumni@2026")
    response = client.get("/api/admin/reports/oaaps", headers=auth_header(alumni["access_token"]))
    assert response.status_code == 403
    assert client.get("/api/admin/reports/oaaps").status_code == 401
