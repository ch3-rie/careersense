from io import BytesIO
from pathlib import Path

import fitz
from docx import Document

from app.services.parser import (
    extract_resume_from_upload,
    extract_resume_info,
    map_to_gts_fields,
    reconcile_extractions,
)
from app.services.parser_document import extract_document, set_ocr_engine_for_tests, text_is_sufficient
from app.services.parser_extract import parse_date_range, parse_date_token, _split_name
from app.services.parser_course import align_occupation_to_degree

STANDARD = """\
Maria Cruz Reyes
Angeles City, Pampanga, Philippines
maria.reyes@gmail.com

EDUCATION
Bachelor of Science in Information Technology
Angeles University Foundation
2018 - 2022

WORK EXPERIENCE
Junior Software Developer
Pampanga Digital Labs | June 2022 - December 2023
Immediate Head: Rico Alvarez (rico.alvarez@nlsc.example)

Software Engineer
North Luzon Systems Corp. | January 2024 - Present
Immediate Head: Rico Alvarez (rico.alvarez@nlsc.example)

SKILLS
Python, SQL, React, Git, REST APIs
"""

NO_HEADINGS = """\
Juan Dela Cruz
juan@example.com
Philippines

BS Information Technology
Angeles University Foundation
2022-2026

Software Developer
ABC Technologies
June 2025 - Present
"""

UNUSUAL_HEADINGS = """\
Sofia Lim
sofia.lim@example.com

Academic Qualifications
Bachelor of Science in Computer Science
AUF 2020

Professional Background
IT Support Specialist — Helix Solutions — Jan 2021 - Dec 2022
Software Engineer — Helix Solutions — 2023-present

Core Competencies
JavaScript, Node.js, SQL
"""

ONE_LINE_JOBS = """\
Paolo Santos
paolo@example.com
Bachelor of Science in Information Technology 2019

Experience
Software Engineer — ABC Technologies — 2022-Present
ABC Technologies • QA Analyst • 2020-2022
"""

MULTIPLE_CURRENT = """\
Lara Tan
lara@example.com
BS Computer Science, AUF, 2021

Work History
Freelance Designer
Studio Nine | 2023 - Present
Part-time Instructor
City College | 2024 - Current
"""

MULTI_DEGREE = """\
Chris Gomez
chris@example.com

Education
Bachelor of Science in Information Technology
Angeles University Foundation
2016-2020

Master of Science in Data Science
State University
2021-2023

Employment
Data Engineer
Insight Labs | 2023 - Present
"""

MISSING_BITS = """\
Nina Cruz
nina@example.com
Bachelor of Arts 2018

Experience
Research Assistant
June 2019 - May 2020
"""

NO_SKILLS = """\
Omar Diaz
omar@example.com
Philippines
Bachelor of Science in Nursing
AUF 2017
Staff Nurse
City Hospital | 2018 - 2022
"""

IRRELEVANT = """\
References available upon request
Curriculum Vitae
Confidential

Mia Flores
mia@example.com
Bachelor of Science in Accountancy
AUF 2019
Junior Auditor
Cruz Accounting Inc. | 2020 - 2021
"""


def _pdf(blocks, width=612, height=792):
    doc = fitz.open()
    page = doc.new_page(width=width, height=height)
    for item in blocks:
        x, y, text = item[0], item[1], item[2]
        size = item[3] if len(item) > 3 else 11
        page.insert_text((x, y), text, fontsize=size)
    data = doc.tobytes()
    doc.close()
    return data


def _docx_table():
    document = Document()
    document.add_paragraph("Ana Santos")
    document.add_paragraph("ana.santos@example.com")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Education"
    table.cell(0, 1).text = "BS Computer Science\nAngeles University Foundation\n2021"
    table.cell(1, 0).text = "Experience"
    table.cell(1, 1).text = "Data Analyst — Metro Analytics — 2021-Present"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _image_only_pdf():
    doc = fitz.open()
    page = doc.new_page()
    page.draw_rect(page.rect, color=(0.8, 0.8, 0.8), fill=(0.8, 0.8, 0.8))
    data = doc.tobytes()
    doc.close()
    return data


class FakeOcr:
    def __init__(self, text):
        self.text = text

    def image_to_text(self, image_bytes):
        return self.text


def test_standard_one_column_resume():
    parsed, source = extract_resume_info(STANDARD)
    assert source == "heuristic"
    assert parsed["first_name"] == "Maria Cruz"
    assert parsed["middle_name"] == ""
    assert parsed["last_name"] == "Reyes"
    assert "Information Technology" in parsed["degree"]
    assert parsed["year_graduated"] == "2022"
    assert parsed["is_currently_employed"] == "Yes"
    assert len(parsed["experiences"]) >= 2
    titles = [job["job_title"] for job in parsed["experiences"]]
    assert any("Software Engineer" in title for title in titles)
    assert "Python" in parsed["skills"]
    gts = map_to_gts_fields(parsed)
    assert gts["pres_occ"]
    assert gts["first_occ"]
    assert gts["present_related_degree"] == "Yes"
    assert parsed["course_alignment"]["status"] == "Aligned"


def test_resume_without_section_headings():
    parsed, _source = extract_resume_info(NO_HEADINGS)
    assert parsed["first_name"] == "Juan"
    assert parsed["last_name"] == "Dela Cruz"
    assert parsed["middle_name"] == ""
    assert "Information Technology" in parsed["degree"]
    assert parsed["experiences"]
    assert parsed["experiences"][-1]["is_current"] == "Yes"
    assert parsed["experiences"][-1]["employer"]


def test_unusual_section_headings():
    parsed, _source = extract_resume_info(UNUSUAL_HEADINGS)
    assert "Computer Science" in parsed["degree"]
    assert any("Software Engineer" in job["job_title"] for job in parsed["experiences"])
    assert any(skill.lower() in {"javascript", "node.js", "sql"} for skill in parsed["skills"])


def test_job_title_employer_date_on_one_line():
    parsed, _source = extract_resume_info(ONE_LINE_JOBS)
    employers = [job.get("employer", "") for job in parsed["experiences"]]
    assert any("ABC Technologies" in employer for employer in employers)
    assert any(job.get("is_current") == "Yes" for job in parsed["experiences"])


def test_separate_line_jobs_and_multiple_roles():
    parsed, _source = extract_resume_info(STANDARD)
    assert len(parsed["experiences"]) >= 2
    current = [job for job in parsed["experiences"] if job["is_current"] == "Yes"]
    assert current
    assert current[0]["end_date"] == ""


def test_multiple_current_jobs():
    parsed, _source = extract_resume_info(MULTIPLE_CURRENT)
    current = [job for job in parsed["experiences"] if job["is_current"] == "Yes"]
    assert len(current) >= 2
    assert parsed["is_currently_employed"] == "Yes"


def test_multiple_degrees_and_further_studies():
    parsed, _source = extract_resume_info(MULTI_DEGREE)
    assert "Information Technology" in parsed["degree"]
    assert parsed["year_graduated"] == "2020"
    assert parsed["further_studies"]
    assert any("Data Science" in row["course_degree"] for row in parsed["further_studies"])
    evidence = parsed["further_studies_evidence"]
    assert evidence["status"] == "detected"
    match = next(row for row in evidence["entries"] if "Data Science" in row["program"])
    assert match["institution"]
    assert match["degree_level"] == "master"
    assert match["confidence"] in {"high", "medium"}


def test_further_studies_evidence_not_detected_for_bachelor_only():
    parsed, _source = extract_resume_info(MISSING_BITS)
    evidence = parsed["further_studies_evidence"]
    assert evidence["status"] == "not_detected"
    assert evidence["entries"] == []


def test_further_studies_evidence_unclear_does_not_invent_a_school():
    from app.services.parser_extract import further_studies_evidence

    evidence = further_studies_evidence([], [], "graduate studies")
    assert evidence["status"] == "unclear"
    assert evidence["entries"][0]["institution"] == ""
    assert evidence["entries"][0]["confidence"] == "low"


def test_date_formats():
    assert parse_date_token("January 2022") == "2022-01"
    assert parse_date_token("Jan 2022") == "2022-01"
    assert parse_date_token("01/2022") == "2022-01"
    assert parse_date_token("01-2022") == "2022-01"
    assert parse_date_token("2022") == "2022"
    start, end, current = parse_date_range("June 2022 - Present")
    assert start.startswith("2022")
    assert end == ""
    assert current == "Yes"
    start, end, current = parse_date_range("2022 - 2024")
    assert start == "2022"
    assert end == "2024"
    assert current == "No"
    start, end, current = parse_date_range("2022-present")
    assert current == "Yes"
    for phrase in ("2022 - Current", "2022 – Now", "Jun 2022 – Dec 2024"):
        start, end, current = parse_date_range(phrase)
        assert start.startswith("2022")
        if "now" in phrase.lower() or "current" in phrase.lower():
            assert current == "Yes"
        else:
            assert end.startswith("2024")


def test_missing_dates_and_employer_stay_blank():
    parsed, _source = extract_resume_info(MISSING_BITS)
    assert parsed["experiences"]
    job = parsed["experiences"][0]
    assert job["job_title"]
    assert job["employer"] == ""
    assert job["start_date"]


def test_no_skills_section_does_not_invent_soft_skills():
    parsed, _source = extract_resume_info(NO_SKILLS)
    lowered = [skill.lower() for skill in parsed["skills"]]
    assert "hospital" not in lowered
    assert "staff" not in lowered


def test_irrelevant_text_is_ignored():
    parsed, _source = extract_resume_info(IRRELEVANT)
    assert parsed["first_name"] == "Mia"
    titles = " ".join(job["job_title"] for job in parsed["experiences"]).lower()
    assert "references" not in titles
    assert "curriculum" not in parsed["first_name"].lower()


def test_sidebar_headings_work_for_other_titles_and_names():
    """A label column is a layout, not a clinical resume. Headings can use other wording."""
    data = _pdf([
        (48, 48, "Engr. Liza B. Santos Jr.", 18),
        (48, 70, "liza.santos@example.com", 9),
        (48, 130, "WORK HISTORY", 11),
        (170, 130, "Civil Engineer"),
        (170, 148, "Sample Builders | 2020 - Present"),
        (48, 210, "ACADEMIC BACKGROUND", 11),
        (170, 210, "Bachelor of Science in Civil Engineering"),
        (170, 228, "Sample University | 2015-2019"),
        (48, 290, "CORE SKILLS", 11),
        (170, 290, "AutoCAD"),
        (400, 48, "Philippines", 9),
    ], width=596, height=842)
    view = extract_document("resume.pdf", data, enable_ocr=False)
    parsed, _source = extract_resume_info(view.text, document=view)
    assert parsed["first_name"] == "Liza"
    assert parsed["middle_name"] == "B."
    assert parsed["last_name"] == "Santos Jr."
    assert "Civil Engineering" in parsed["degree"]
    assert parsed["year_graduated"] == "2019"
    titles = [job["job_title"] for job in parsed["experiences"]]
    assert any("Civil Engineer" in title for title in titles)
    assert "AutoCAD" not in titles
    assert "AutoCAD" in parsed["skills"]


def test_right_side_labels_stay_with_their_sections():
    data = _pdf([
        (48, 48, "Paolo Reyes", 18),
        (48, 70, "paolo.reyes@example.com", 9),
        (48, 140, "Account Officer"),
        (48, 158, "Sample Credit Co. | 2023 - Present"),
        (48, 220, "Bachelor of Science in Accountancy"),
        (48, 238, "Sample University 2022"),
        (400, 140, "EMPLOYMENT", 11),
        (400, 220, "EDUCATION", 11),
        (400, 300, "SKILLS", 11),
        (48, 300, "Excel"),
    ], width=596, height=842)
    view = extract_document("resume.pdf", data, enable_ocr=False)
    parsed, _source = extract_resume_info(view.text, document=view)
    assert parsed["first_name"] == "Paolo"
    assert parsed["last_name"] == "Reyes"
    assert "Accountancy" in parsed["degree"]
    assert any("Account Officer" in job["job_title"] for job in parsed["experiences"])
    assert "Excel" in parsed["skills"]


def test_sidebar_headings_stay_with_the_name_and_job():
    data = _pdf([
        (45, 50, "Dr. Marcus Vance, Ph.D.", 20),
        (45, 72, "Lead Clinical Research Scientist"),
        (430, 50, "m.vance@bioresearch.org", 9),
        (45, 130, "PROFILE", 11),
        (160, 130, "Senior researcher in clinical trials."),
        (45, 180, "EXPERIENCE", 11),
        (160, 180, "Principal Clinical Investigator"),
        (160, 196, "St. Luke's Medical Center | 2022 - Present"),
        (45, 260, "EDUCATION", 11),
        (160, 260, "B.S. in Biochemistry"),
        (160, 276, "Sample University 2015"),
        (45, 340, "PUBLICATIONS", 11),
        (160, 340, "Vance, M. (2024). Targeted therapies."),
        (45, 400, "SKILLS", 11),
        (160, 400, "Flow Cytometry"),
    ], width=596, height=842)
    view = extract_document("resume.pdf", data, enable_ocr=False)
    parsed, _source = extract_resume_info(view.text, document=view)
    assert parsed["first_name"] == "Marcus"
    assert parsed["last_name"] == "Vance"
    titles = [job["job_title"] for job in parsed["experiences"]]
    assert any("Principal Clinical Investigator" in title for title in titles)
    assert "Flow Cytometry" not in titles
    assert "Biochemistry" in parsed["degree"]


def test_two_column_pdf_keeps_column_order():
    data = _pdf([
        (72, 50, "Diego Ramos", 16),
        (72, 68, "diego@example.com"),
        (50, 110, "EDUCATION"),
        (50, 126, "BSIT"),
        (50, 142, "Angeles University Foundation 2022"),
        (50, 180, "EXPERIENCE"),
        (50, 196, "Software Engineer"),
        (50, 212, "ABC Technologies | 2022 - Present"),
        (340, 110, "SKILLS"),
        (340, 126, "Python"),
        (340, 142, "SQL"),
        (340, 180, "CERTIFICATIONS"),
        (340, 196, "AWS Cloud Practitioner"),
    ])
    view = extract_document("resume.pdf", data, enable_ocr=False)
    parsed, _source = extract_resume_info(view.text, document=view)
    assert parsed["first_name"] == "Diego"
    blob = " ".join(job["job_title"] for job in parsed["experiences"]).lower()
    assert "python" not in blob
    assert any("Software Engineer" in job["job_title"] for job in parsed["experiences"])


def test_llamaparse_text_is_used_before_the_local_reader(monkeypatch):
    calls = []

    def _llama(filename, file_bytes):
        calls.append(filename)
        return (
            "Ana Cruz\n"
            "Bachelor of Science in Information Technology\n"
            "Software Engineer\nNorth Labs\n2022-Present\n"
            "Skills: Python\n"
        )

    def _local(*_args, **_kwargs):
        raise AssertionError("local reader should wait until LlamaParse returns nothing")

    monkeypatch.setattr("app.services.parser.parse_with_llamaparse", _llama)
    monkeypatch.setattr("app.services.parser.extract_document", _local)
    parsed, source, text = extract_resume_from_upload("resume.pdf", b"%PDF-1.4")
    assert calls == ["resume.pdf"]
    assert "Ana" in parsed["first_name"]
    assert "Ana Cruz" in text
    assert source == "llamaparse+heuristic"


def test_table_based_docx():
    parsed, source, text = extract_resume_from_upload("resume.docx", _docx_table())
    assert "Ana" in parsed["first_name"]
    assert text_is_sufficient(text)
    assert parsed["experiences"] or "analyst" in text.lower()
    assert source in {"heuristic", "llamaindex+heuristic"}


def test_scanned_pdf_without_ocr_is_empty():
    view = extract_document("scan.pdf", _image_only_pdf(), enable_ocr=False)
    parsed, source = extract_resume_info(view.text, document=view)
    assert source in {"empty", "failed"}
    assert not parsed.get("first_name")
    assert not text_is_sufficient(view.text)


def test_scanned_pdf_with_ocr_fallback():
    set_ocr_engine_for_tests(FakeOcr(STANDARD))
    try:
        parsed, source, text = extract_resume_from_upload("scan.pdf", _image_only_pdf())
        assert source == "ocr"
        assert text_is_sufficient(text)
        assert parsed["last_name"] == "Reyes"
        assert parsed["experiences"]
    finally:
        set_ocr_engine_for_tests(None)


def test_empty_and_unreadable_file():
    parsed, source, text = extract_resume_from_upload("blank.txt", b"   \n")
    assert source == "empty"
    assert parsed == {}
    assert text == ""
    parsed, source = extract_resume_info("")
    assert source == "empty"


def test_long_content_is_truncated_safely():
    padding = ("lorem ipsum dolor sit amet " * 80 + "\n") * 40
    parsed, source = extract_resume_info(STANDARD + "\n" + padding)
    assert source == "heuristic"
    assert parsed["last_name"] == "Reyes"


def test_conflicting_ai_and_heuristic_prefers_document_evidence():
    heuristic = {
        "first_name": "Maria",
        "last_name": "Reyes",
        "degree": "Bachelor of Science in Information Technology",
        "experiences": [{"job_title": "IT Support Specialist", "employer": "Helix Solutions", "start_date": "2021-01"}],
        "skills": ["Python"],
        "further_studies": [],
        "ever_employed": "Yes",
        "is_currently_employed": "",
    }
    ai = {
        "first_name": "Maria",
        "last_name": "Reyes",
        "degree": "Doctor of Philosophy",
        "experiences": [{"job_title": "Software Engineer", "employer": "Invented Corp", "start_date": "2021-01"}],
        "skills": ["Teleportation"],
        "further_studies": [],
    }
    merged, confidence = reconcile_extractions(ai, heuristic, STANDARD)
    assert "Information Technology" in merged["degree"]
    titles = [job["job_title"] for job in merged["experiences"]]
    assert "IT Support Specialist" in titles
    assert "Invented Corp" not in [job.get("employer") for job in merged["experiences"]]
    assert "Software Engineer" not in titles
    assert "Teleportation" not in merged["skills"]
    assert confidence["degree"] in {"medium", "high", "low"}


def test_bundled_sample_resume():
    path = Path(__file__).resolve().parents[1] / "data" / "sample_resume.txt"
    parsed, source = extract_resume_info(path.read_text(encoding="utf-8"))
    assert source == "heuristic"
    assert parsed["first_name"] == "Maria Cruz"
    assert parsed["middle_name"] == ""
    assert parsed["last_name"] == "Reyes"
    assert "Information Technology" in parsed["degree"]
    assert parsed["year_graduated"] == "2022"
    assert parsed["is_currently_employed"] == "Yes"
    assert len(parsed["experiences"]) >= 2
    assert "Python" in parsed["skills"]


def test_same_employer_different_titles_are_not_first_job():
    gts = map_to_gts_fields(
        {
            "experiences": [
                {"job_title": "Software Developer", "employer": "ABC Corporation", "is_current": "No"},
                {"job_title": "Systems Analyst", "employer": "ABC Corporation", "is_current": "Yes"},
            ],
            "is_currently_employed": "Yes",
        }
    )
    assert gts["present_job_is_first"] == "No"
    assert gts["first_occ"] == "Software Developer"
    assert gts["pres_occ"] == "Systems Analyst"


def test_empty_extraction_maps_to_blank_gts_fields():
    gts = map_to_gts_fields({})
    assert gts["first_name"] == ""
    assert gts["pres_occ"] == ""
    assert gts["experiences"] == []
    assert gts["skills"] == []


def test_missing_current_job_is_not_assumed():
    parsed, _source = extract_resume_info(NO_SKILLS)
    assert parsed["is_currently_employed"] == "No"
    gts = map_to_gts_fields(parsed)
    assert gts["pres_occ"] == ""


def test_digital_pdf_does_not_require_ocr():
    data = _pdf([
        (72, 72, "Maria Cruz Reyes", 16),
        (72, 94, "Bachelor of Science in Information Technology"),
        (72, 110, "Angeles University Foundation 2022"),
        (72, 140, "Software Engineer"),
        (72, 156, "North Luzon Systems Corp. | January 2024 - Present"),
        (72, 190, "Skills: Python, SQL"),
    ])
    view = extract_document("resume.pdf", data, enable_ocr=False)
    assert view.source == "digital_pdf"
    parsed, source = extract_resume_info(view.text, document=view)
    assert source == "heuristic"
    assert parsed["first_name"] == "Maria Cruz"


def test_three_word_name_keeps_compound_given_name():
    parsed = _split_name("Sean Gabriel Santos")
    assert parsed["first_name"] == "Sean Gabriel"
    assert parsed["middle_name"] == ""
    assert parsed["last_name"] == "Santos"


def test_four_word_filipino_name_can_split_middle():
    parsed = _split_name("Maria Clara Reyes Santos")
    assert parsed["first_name"] == "Maria Clara"
    assert parsed["middle_name"] == "Reyes"
    assert parsed["last_name"] == "Santos"


def test_compound_surname_dela_cruz():
    parsed = _split_name("Juan Dela Cruz")
    assert parsed["first_name"] == "Juan"
    assert parsed["middle_name"] == ""
    assert parsed["last_name"] == "Dela Cruz"


def test_explicit_name_labels_take_priority():
    resume = """\
First Name: Sean Gabriel
Middle Name: Cruz
Last Name: Santos

Bachelor of Science in Information Technology
AUF 2022
Software Engineer
Helix Labs | 2023 - Present
"""
    parsed, _source = extract_resume_info(resume)
    assert parsed["first_name"] == "Sean Gabriel"
    assert parsed["middle_name"] == "Cruz"
    assert parsed["last_name"] == "Santos"


def test_full_name_label_does_not_invent_middle_name():
    resume = """\
Name: Sean Gabriel Santos
sean.gabriel@example.com
Bachelor of Science in Computer Science
AUF 2021
Web Developer
Pixel Forge | 2022 - Current
"""
    parsed, _source = extract_resume_info(resume)
    assert parsed["first_name"] == "Sean Gabriel"
    assert parsed["middle_name"] == ""
    assert parsed["last_name"] == "Santos"


def test_current_job_alignment_not_inherited_from_previous_role():
    resume = """\
Sean Gabriel Santos
sean@example.com
Bachelor of Science in Information Technology
Angeles University Foundation 2022

Work Experience
Software Developer
Pampanga Digital Labs | 2022 - 2023

Marketing Specialist
North Star Media | 2024 - Present
"""
    parsed, _source = extract_resume_info(resume)
    current = [job for job in parsed["experiences"] if job["is_current"] == "Yes"]
    assert current
    assert "Marketing" in current[-1]["job_title"]
    assert parsed["course_alignment"]["status"] == "Misaligned"
    assert parsed["course_alignment"]["job_title"]
    gts = map_to_gts_fields(parsed)
    assert gts["present_related_degree"] == "No"
    assert gts["pres_occ"]
    assert "Marketing" in gts["pres_occ"]


def test_it_roles_align_without_exact_degree_words():
    degree = "Bachelor of Science in Information Technology"
    for title in ("QA Engineer", "Web Developer", "Business Analyst", "IT Support Specialist"):
        result = align_occupation_to_degree(degree, title)
        assert result.status == "Aligned", title
        assert result.related_to_degree == "Yes"
        assert result.confidence >= 0.7


def test_unrelated_current_role_is_misaligned():
    result = align_occupation_to_degree(
        "Bachelor of Science in Information Technology",
        "Registered Nurse",
    )
    assert result.status == "Misaligned"
    assert result.related_to_degree == "No"


def test_ambiguous_title_without_evidence_is_unknown():
    result = align_occupation_to_degree(
        "Bachelor of Science in Information Technology",
        "Analyst",
    )
    assert result.status == "Unknown"
    assert result.related_to_degree == ""


def test_ambiguous_title_uses_responsibilities():
    unknown = align_occupation_to_degree(
        "Bachelor of Science in Information Technology",
        "Coordinator",
    )
    assert unknown.status == "Unknown"
    result = align_occupation_to_degree(
        "Bachelor of Science in Information Technology",
        "Coordinator",
        description="Developed web applications, maintained databases, and configured network infrastructure.",
    )
    assert result.status == "Aligned"
    assert result.related_to_degree == "Yes"


def test_no_current_job_course_alignment_is_unknown():
    parsed, _source = extract_resume_info(NO_SKILLS)
    assert parsed["is_currently_employed"] == "No"
    assert parsed["course_alignment"]["status"] == "Unknown"
    assert parsed["course_alignment"]["related_to_degree"] == ""


def test_page_number_is_not_a_certification_and_credential_fields_match_the_profile():
    resume = """\
Ana Santos
ana@example.com
Bachelor of Science in Information Technology
Angeles University Foundation
2018-2022

Certifications
Google Analytics Certificate | Google | 2023
Page 1 of 1

Seminars Attended
Career Readiness Seminar | OAAPS | Angeles City | 2022

Awards and Achievements
Dean's Lister | AUF | College honors
"""
    parsed, _source = extract_resume_info(resume)
    assert all("page" not in item["name"].lower() for item in parsed["certifications"])
    certificate = next(item for item in parsed["certifications"] if "Google Analytics" in item["name"])
    assert certificate["awarding_body"] == "Google"
    assert certificate["year"] == "2023"
    assert certificate["detail"] != "2023"
    seminar = parsed["seminars"][0]
    assert seminar["organizer"] == "OAAPS"
    assert seminar["location"] == "Angeles City"
    assert seminar["year"] == "2022"
    award = parsed["awards"][0]
    assert award["issuer"] == "AUF"
    assert award["description"] == "College honors"
    assert award["year"] == ""


def test_survey_uses_work_after_graduation_and_keeps_credentials():
    resume = """\
Ana Santos
ana@example.com
Bachelor of Science in Information Technology
Angeles University Foundation
2018-2022

Work Experience
Marketing Intern
Sample Office
January 2017 - March 2018

Software Engineer
North Labs
July 2022 - Present

Certifications
Google Analytics Certificate | 2023

Seminars Attended
Career Readiness Seminar | Angeles City

Awards and Achievements
Dean's Lister | AUF
"""
    parsed, _source = extract_resume_info(resume)
    gts = map_to_gts_fields(parsed)
    assert gts["ever_employed"] == "Yes"
    assert gts["first_occ"] == "Software Engineer"
    assert gts["first_emp"] == "North Labs"
    assert gts["time_to_first_job"] == "1 to 6 months"
    assert gts["present_related_degree"] in {"Yes", "No"}
    assert "Information Technology" in gts["degree"]
    assert any("Intern" in job["job_title"] for job in gts["experiences"])
    assert any("Google Analytics" in item["name"] for item in gts["certifications"])
    assert any("Career Readiness" in item["name"] for item in gts["seminars"])
    assert any("Dean" in item["name"] for item in gts["awards"])


def test_same_year_without_a_month_counts_as_one_to_six_months():
    gts = map_to_gts_fields({
        "year_graduated": "2019",
        "degree": "BS Business Administration",
        "experiences": [
            {"job_title": "Specialist", "employer": "Agency", "start_date": "2019", "is_current": "No", "related_to_degree": "Yes"},
            {"job_title": "Lead", "employer": "Media", "start_date": "2023-01", "is_current": "Yes", "related_to_degree": "Yes"},
        ],
    })
    assert gts["ever_employed"] == "Yes"
    assert gts["first_occ"] == "Specialist"
    assert gts["time_to_first_job"] == "1 to 6 months"
    assert gts["pres_occ"] == "Lead"


def test_marketing_jobs_answer_relatedness_and_alignment():
    from app.services.parser import normalize_extracted_resume_data

    parsed = normalize_extracted_resume_data({
        "degree": "BS Business Administration Major in Marketing",
        "year_graduated": "2019",
        "experiences": [
            {
                "job_title": "SEO & Content Specialist",
                "employer": "Vanguard Creative Agency",
                "start_date": "2019",
                "end_date": "2021",
                "is_current": "No",
                "description": "Managed technical SEO audits, keyword mapping, and content campaigns.",
            },
            {
                "job_title": "Growth Marketing Strategist",
                "employer": "Nexus Digital Labs",
                "start_date": "2021",
                "end_date": "2023",
                "is_current": "No",
            },
            {
                "job_title": "Senior Digital Marketing Lead",
                "employer": "Apex Media Group Philippines",
                "start_date": "2023",
                "is_current": "Yes",
                "description": "Led digital marketing campaigns across paid ads and SEO.",
            },
        ],
        "skills": ["SEO", "Google Ads", "Copywriting"],
    }, "")
    gts = map_to_gts_fields(parsed)
    assert gts["first_related"] == "Yes"
    assert gts["present_related_degree"] == "Yes"
    assert gts["time_to_first_job"] == "1 to 6 months"
    assert parsed["course_alignment"]["status"] == "Aligned"
    assert parsed["course_alignment"]["related_to_degree"] == "Yes"


def test_work_that_ended_before_graduation_is_not_employment_after_graduation():
    gts = map_to_gts_fields({
        "year_graduated": "2022",
        "experiences": [
            {"job_title": "Intern", "employer": "Shop", "start_date": "2017-01", "end_date": "2018-03", "is_current": "No"},
        ],
    })
    assert gts["ever_employed"] == "No"
    assert gts["first_occ"] == ""
    assert gts["time_to_first_job"] == ""


def test_resume_history_and_credentials_reach_the_alumni_home(client, db_session):
    from app.models import Account, AlumniAward, AlumniCertification, AlumniProfile, AlumniSeminar, Resume, TracerSubmission
    from app.services.jobs import serialize_jobs, sync_job_timeline_from_gts
    from app.services.profile_records import seed_resume_activity

    account = Account(
        personal_email="timeline-seed@example.com",
        password_hash="x",
        role="Alumni",
        status="Pending",
        privacy_consent=True,
    )
    db_session.add(account)
    db_session.flush()
    db_session.add(AlumniProfile(
        account_id=account.id,
        first_name="Ana",
        last_name="Santos",
        degree="BS Information Technology",
        year_graduated="2022",
    ))
    db_session.add(Resume(
        account_id=account.id,
        original_filename="resume.pdf",
        stored_path="uploads/resume.pdf",
        parsed_json={
            "experiences": [
                {"job_title": "Intern", "employer": "Shop", "start_date": "2017-01", "end_date": "2018-03", "is_current": "No"},
                {"job_title": "Engineer", "employer": "North Labs", "start_date": "2022-07", "end_date": "2023-06", "is_current": "No"},
                {"job_title": "Lead", "employer": "Media", "start_date": "2023-07", "is_current": "Yes"},
            ],
            "certifications": [{"name": "Google Analytics Certificate", "detail": "Google", "year": "2023"}],
            "seminars": [{"name": "Career Readiness Seminar", "detail": "Angeles City", "year": "2022"}],
            "awards": [{"name": "Dean's Lister", "detail": "AUF", "year": "2021"}],
        },
    ))
    db_session.add(TracerSubmission(
        account_id=account.id,
        data_json={"first_occ": "Engineer", "first_emp": "North Labs", "pres_occ": "Lead", "is_currently_employed": "Yes"},
        extra_answers={},
    ))
    db_session.flush()
    db_session.refresh(account)
    sync_job_timeline_from_gts(db_session, account, {"first_occ": "Engineer", "pres_occ": "Lead", "is_currently_employed": "Yes", "present_job_is_first": "No"})
    db_session.flush()
    titles = [job["job_title"] for job in serialize_jobs(db_session, account.id)]
    assert titles == ["Intern", "Engineer", "Lead"]
    seed_resume_activity(db_session, account)
    db_session.commit()
    certificate = db_session.query(AlumniCertification).filter_by(account_id=account.id).one()
    seminar = db_session.query(AlumniSeminar).filter_by(account_id=account.id).one()
    award = db_session.query(AlumniAward).filter_by(account_id=account.id).one()
    assert certificate.issuing_organization == "Google"
    assert certificate.awarded_on == "2023"
    assert seminar.location == "Angeles City"
    assert seminar.held_on == "2022"
    assert award.issuer == "AUF"
    assert award.awarded_on == "2021"
    seed_resume_activity(db_session, account)
    assert db_session.query(AlumniCertification).filter_by(account_id=account.id).count() == 1
