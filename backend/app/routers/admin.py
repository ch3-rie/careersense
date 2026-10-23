import csv
import io
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import load_workbook
from sqlalchemy import String, and_, cast, func, literal, or_
from sqlalchemy.orm import Session, joinedload

from app.constants import CORE_GTS_SECTIONS, QUESTION_TYPES
from app.db import get_db
from app.deps import require_admin
from app.models import (
    Account,
    AdminLog,
    AlignmentResult,
    AlumniCard,
    AlumniProfile,
    FurtherStudy,
    JobTitleMapping,
    Resume,
    SocCode,
    TracerSubmission,
    UniversityRecord,
)
from app.routers.auth import _norm, match_university_record
from app.schemas import GtsPayload, MappingIn, QuestionIn, QuestionUpdate, RejectRequest, SocIn, UniversityRecordIn
from app.services.admin_directory import (
    directory_export_rows,
    directory_filename,
    directory_page,
    render_directory_csv,
    render_directory_xlsx,
)
from app.services.academic_filters import (
    ProgramSelectionError,
    academic_catalog,
    academic_programs,
    apply_registry_academic_filters,
    apply_tracer_academic_filters,
    ensure_academic_catalog,
    resolve_program_selection,
)
from app.services.academic_catalog import course_code_owner, sync_catalog_from_records
from app.services.registry_import import (
    RegistryImportError,
    commit_registry_import,
    latest_registry_import,
    preview_registry_import,
)
from app.services.alignment import invalidate_soc_cache
from app.services.email import send_account_status_email
from app.services.gts import persist_tracer_submission
from app.services.survey import get_published_schema, validate_required_answers
from app.services.employment import apply_profile_unemployment, resolve_current_employment
from app.services.files import resolve_stored_file, safe_download_name
from app.services.profile_records import certification_file, serialize_activity
from app.services.studies import serialize_studies
from app.services.oaaps_docx import render_oaaps_docx
from app.services.reports import (
    _latest_tracer_rows,
    _relatedness,
    build_oaaps_report,
    filter_options,
    employment_analytics,
    latest_submissions,
    parse_report_date,
    prepare_employment_row,
)
from app.services.survey import add_legacy_extra, delete_legacy_extra, ensure_gts_survey, extras_payload, save_draft
from app.services.validation import csv_safe, student_id_in_use
from app.survey_default import clone_schema, iter_questions

router = APIRouter(prefix="/api/admin", tags=["admin"])


def _log(db: Session, admin_id: int, action: str, target_id, old=None, new=None):
    db.add(
        AdminLog(
            admin_id=admin_id,
            action_type=action,
            target_id=str(target_id),
            old_value=old,
            new_value=new,
        )
    )


def _account_name(account: Account) -> str:
    if account.profile:
        parts = [account.profile.first_name, account.profile.middle_name, account.profile.last_name]
        name = " ".join(p for p in parts if p).strip()
        if name:
            return name
    return account.personal_email


@router.get("/dashboard")
def dashboard(admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    alumni_q = db.query(Account).filter(Account.role == "Alumni")
    total_registered = alumni_q.count()
    pending = alumni_q.filter(Account.status == "Pending").count()
    approved = alumni_q.filter(Account.status == "Active").count()
    rejected = alumni_q.filter(Account.status == "Rejected").count()
    submissions = db.query(TracerSubmission).count()
    aligned = db.query(AlignmentResult).filter(AlignmentResult.alignment_justification == "Aligned").count()
    unknown = db.query(AlignmentResult).filter(AlignmentResult.alignment_justification == "Unknown").count()
    misaligned = db.query(AlignmentResult).filter(AlignmentResult.alignment_justification == "Misaligned").count()
    pending_cards = (
        db.query(AlumniCard)
        .join(Account, Account.id == AlumniCard.account_id)
        .filter(Account.role == "Alumni", AlumniCard.status == "ForVerification")
        .count()
    )
    classified = aligned + unknown + misaligned
    percent_aligned = round((aligned / classified) * 100, 1) if classified else 0

    trend_rows = (
        db.query(func.date(TracerSubmission.submitted_at), func.count(TracerSubmission.id))
        .group_by(func.date(TracerSubmission.submitted_at))
        .order_by(func.date(TracerSubmission.submitted_at))
        .all()
    )
    return {
        "pending_approvals": pending,
        "pending_card_applications": pending_cards,
        "total_submissions": submissions,
        "approved_alumni": approved,
        "rejected_alumni": rejected,
        "total_registered": total_registered,
        "percent_aligned": percent_aligned,
        "avg_alignment": percent_aligned,
        "alignment_distribution": {"Aligned": aligned, "Unknown": unknown, "Misaligned": misaligned},
        "submission_trend": [{"date": str(d), "count": c} for d, c in trend_rows if d],
    }


@router.get("/directory")
def alumni_directory(
    cohort: str = "all",
    status: str = "",
    alignment: str = "",
    year: str = "",
    answer: str = "",
    reason: str = "",
    college: str = "",
    program: str = "",
    indicator: str = "",
    period_from: str = "",
    period_to: str = "",
    location: str = "",
    employment: str = "",
    award: str = "",
    further: str = "",
    page: int = 1,
    page_size: int = 12,
    q: str = "",
    sort_by: str = "name",
    sort_dir: str = "asc",
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    return directory_page(
        db,
        page=page,
        page_size=page_size,
        q=q,
        sort_by=sort_by,
        sort_dir=sort_dir,
        cohort=cohort,
        status=status,
        alignment=alignment,
        year=year,
        answer=answer,
        reason=reason,
        college=college,
        program=program,
        indicator=indicator,
        period_from=period_from,
        period_to=period_to,
        location=location,
        employment=employment,
        award=award,
        further=further,
    )


@router.get("/directory/export")
def export_directory(
    format: str = "csv",
    cohort: str = "all",
    status: str = "",
    alignment: str = "",
    year: str = "",
    answer: str = "",
    reason: str = "",
    college: str = "",
    program: str = "",
    indicator: str = "",
    period_from: str = "",
    period_to: str = "",
    location: str = "",
    employment: str = "",
    award: str = "",
    further: str = "",
    q: str = "",
    sort_by: str = "name",
    sort_dir: str = "asc",
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    kind = (format or "csv").strip().lower()
    if kind not in {"csv", "xlsx"}:
        raise HTTPException(status_code=400, detail="Choose a CSV or Excel download.")
    title, rows = directory_export_rows(
        db,
        q=q,
        sort_by=sort_by,
        sort_dir=sort_dir,
        cohort=cohort,
        status=status,
        alignment=alignment,
        year=year,
        answer=answer,
        reason=reason,
        college=college,
        program=program,
        indicator=indicator,
        period_from=period_from,
        period_to=period_to,
        location=location,
        employment=employment,
        award=award,
        further=further,
    )
    filename = directory_filename(title, "csv" if kind == "csv" else "xlsx")
    if kind == "csv":
        return StreamingResponse(
            iter([render_directory_csv(rows)]),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    return StreamingResponse(
        iter([render_directory_xlsx(rows)]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


APPROVAL_STATUSES = ("Pending", "Active", "Rejected")


def _course_colleges(db: Session) -> dict[str, list[str]]:
    mapping: dict[str, list[str]] = {}
    for item in academic_programs(db):
        colleges = mapping.setdefault(item["course"].casefold(), [])
        if item["college"] not in colleges:
            colleges.append(item["college"])
    return mapping


def _approval_college(user: Account, course_colleges: dict[str, list[str]]) -> str:
    record = user.university_record
    recorded = (record.college or "").strip() if record else ""
    if recorded:
        return recorded
    degree = ((user.profile.degree if user.profile else "") or "").strip()
    matches = course_colleges.get(degree.casefold(), [])
    return matches[0] if len(matches) == 1 else ""


def _approval_years(db: Session) -> list[str]:
    profile_years = (
        db.query(AlumniProfile.year_graduated)
        .join(Account, Account.id == AlumniProfile.account_id)
        .filter(Account.role == "Alumni")
        .distinct()
        .all()
    )
    record_years = db.query(UniversityRecord.year_graduated).distinct().all()
    years = set()
    for (value,) in [*profile_years, *record_years]:
        text = str(value or "").strip()
        if text:
            years.add(text)
    return sorted(years, reverse=True)


@router.get("/approvals")
def approvals(
    page: int = 1,
    page_size: int = 10,
    q: str = "",
    college: str = "",
    course: str = "",
    year: str = "",
    status: str = "Pending",
    sort_by: str = "created",
    sort_dir: str = "desc",
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    page = max(1, page)
    page_size = min(50, max(5, page_size))
    course_colleges = _course_colleges(db)
    query = (
        db.query(Account)
        .outerjoin(AlumniProfile, AlumniProfile.account_id == Account.id)
        .outerjoin(UniversityRecord, UniversityRecord.student_id == Account.linked_student_id)
        .options(joinedload(Account.profile), joinedload(Account.university_record))
        .filter(Account.role == "Alumni")
    )
    status_aliases = {"approved": "Active", "active": "Active", "pending": "Pending", "rejected": "Rejected", "all": "All"}
    requested_status = status_aliases.get((status or "Pending").strip().casefold(), "Pending")
    term = (q or "").strip()
    if term:
        like = f"%{term}%"
        query = query.filter(
            or_(
                Account.personal_email.ilike(like),
                Account.linked_student_id.ilike(like),
                AlumniProfile.first_name.ilike(like),
                AlumniProfile.middle_name.ilike(like),
                AlumniProfile.last_name.ilike(like),
            )
        )
    college_name = (college or "").strip()
    if college_name:
        record_college = func.lower(func.trim(UniversityRecord.college))
        degree_key = func.lower(func.trim(AlumniProfile.degree))
        unique_courses = [
            name
            for name, colleges in course_colleges.items()
            if len(colleges) == 1 and colleges[0].casefold() == college_name.casefold()
        ]
        query = query.filter(
            or_(
                record_college == college_name.casefold(),
                and_(
                    or_(UniversityRecord.college.is_(None), func.trim(UniversityRecord.college) == ""),
                    degree_key.in_(unique_courses or ["__none__"]),
                ),
            )
        )
    course_name = (course or "").strip()
    if course_name:
        query = query.filter(
            or_(
                func.lower(func.trim(AlumniProfile.degree)) == course_name.casefold(),
                func.lower(func.trim(UniversityRecord.degree)) == course_name.casefold(),
            )
        )
    batch_year = (year or "").strip()
    if batch_year:
        query = query.filter(
            or_(
                func.trim(AlumniProfile.year_graduated) == batch_year,
                and_(
                    or_(AlumniProfile.year_graduated.is_(None), func.trim(AlumniProfile.year_graduated) == ""),
                    func.trim(UniversityRecord.year_graduated) == batch_year,
                ),
            )
        )
    summary_rows = (
        query.order_by(None)
        .with_entities(Account.status, func.count(func.distinct(Account.id)))
        .group_by(Account.status)
        .all()
    )
    summary_counts = {row_status: count for row_status, count in summary_rows}
    summary = {
        "approved": summary_counts.get("Active", 0),
        "pending": summary_counts.get("Pending", 0),
        "rejected": summary_counts.get("Rejected", 0),
    }
    if requested_status != "All":
        query = query.filter(Account.status == requested_status)
    key = (sort_by or "created").lower()
    descending = (sort_dir or "desc").lower() != "asc"
    if key == "email":
        order_col = func.lower(Account.personal_email)
    elif key == "name":
        order_col = func.lower(func.coalesce(AlumniProfile.last_name, Account.personal_email))
    else:
        order_col = Account.created_at
    query = query.order_by(order_col.desc() if descending else order_col.asc(), Account.id.desc())
    total = query.count()
    rows = query.offset((page - 1) * page_size).limit(page_size).all()
    items = []
    for user in rows:
        latest = (
            db.query(TracerSubmission)
            .filter(TracerSubmission.account_id == user.id)
            .order_by(TracerSubmission.submitted_at.desc())
            .first()
        )
        resume = (
            db.query(Resume)
            .filter(Resume.account_id == user.id)
            .order_by(Resume.created_at.desc())
            .first()
        )
        profile_degree = (user.profile.degree if user.profile else "") or ""
        record_degree = (user.university_record.degree if user.university_record else "") or ""
        profile_year = (user.profile.year_graduated if user.profile else "") or ""
        record_year = (user.university_record.year_graduated if user.university_record else "") or ""
        items.append({
            "id": user.id,
            "email": user.personal_email,
            "student_id": user.linked_student_id,
            "status": user.status,
            "created_at": user.created_at,
            "name": _account_name(user),
            "college": _approval_college(user, course_colleges),
            "degree": profile_degree.strip() or record_degree.strip(),
            "year_graduated": profile_year.strip() or record_year.strip(),
            "matched": bool(user.linked_student_id),
            "has_submission": bool(latest),
            "resume_id": resume.id if resume else None,
            "resume_filename": resume.original_filename if resume else None,
        })
    return {
        "items": items,
        "total": total,
        "page": page,
        "page_size": page_size,
        "q": term,
        "college": college_name,
        "course": course_name,
        "year": batch_year,
        "status": "Approved" if requested_status == "Active" else requested_status,
        "summary": summary,
        "years": _approval_years(db),
        "statuses": ["All", *APPROVAL_STATUSES],
        "sort_by": key if key in {"name", "email", "created"} else "created",
        "sort_dir": "desc" if descending else "asc",
    }


@router.get("/approvals/{account_id}")
def approval_detail(account_id: int, admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    user = db.get(Account, account_id)
    if not user or user.role != "Alumni":
        raise HTTPException(status_code=404, detail="Alumni account not found.")
    latest = (
        db.query(TracerSubmission)
        .filter(TracerSubmission.account_id == user.id)
        .order_by(TracerSubmission.submitted_at.desc())
        .first()
    )
    studies = db.query(FurtherStudy).filter(FurtherStudy.account_id == user.id).all()
    resumes = db.query(Resume).filter(Resume.account_id == user.id).order_by(Resume.created_at.desc()).all()
    return {
        "account": {
            "id": user.id,
            "email": user.personal_email,
            "status": user.status,
            "student_id": user.linked_student_id,
            "created_at": user.created_at,
            "is_verified": user.is_verified,
        },
        "profile": {
            "first_name": user.profile.first_name if user.profile else "",
            "middle_name": user.profile.middle_name if user.profile else "",
            "last_name": user.profile.last_name if user.profile else "",
            "husband_surname": user.profile.husband_surname if user.profile else "",
            "country_residence": user.profile.country_residence if user.profile else "",
            "degree": user.profile.degree if user.profile else "",
            "year_graduated": user.profile.year_graduated if user.profile else "",
            "guardian_type": user.profile.guardian_type if user.profile else "",
        } if user.profile else None,
        "university_record": {
            "student_id": user.university_record.student_id,
            "first_name": user.university_record.first_name,
            "middle_name": user.university_record.middle_name,
            "last_name": user.university_record.last_name,
            "personal_email": user.university_record.personal_email or "",
            "degree": user.university_record.degree,
            "year_graduated": user.university_record.year_graduated,
            "college": user.university_record.college,
        } if user.university_record else None,
        "submission": {
            "id": latest.id,
            "submitted_at": latest.submitted_at,
            "alignment_status": latest.alignment_status,
            "data": latest.data_json,
            "extra_answers": latest.extra_answers,
        }
        if latest
        else None,
        "further_studies": [
            {
                "id": s.id,
                "course_degree": s.course_degree,
                "school": s.school,
                "institution_address": s.institution_address or "",
                "year_enrolled": s.year_enrolled,
                "scholarship": s.scholarship,
                "is_graduated": s.is_graduated,
            }
            for s in studies
        ],
        "profile_activity": serialize_activity(db, user.id),
        "resumes": [
            {"id": r.id, "filename": r.original_filename, "created_at": r.created_at, "parsed": r.parsed_json}
            for r in resumes
        ],
    }


@router.post("/approvals/{account_id}/verify")
def verify_graduate(account_id: int, admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    user = db.get(Account, account_id)
    if not user:
        raise HTTPException(status_code=404, detail="Account not found.")
    profile = user.profile
    record = None
    reason = "No graduate record match found"
    matched = False
    match_type = "none"
    if user.linked_student_id:
        record = db.get(UniversityRecord, user.linked_student_id)
        if record:
            matched, reason, match_type = True, f"Matched by Student ID ({record.student_id})", "student_id"
    if not matched:
        record = match_university_record(
            db,
            user.personal_email,
            profile.first_name if profile else "",
            profile.last_name if profile else "",
            profile.middle_name if profile else "",
        )
        if record:
            if student_id_in_use(db, record.student_id, exclude_account_id=user.id):
                return {
                    "matched": False,
                    "match_type": "already_linked",
                    "reason": "This graduate record is already linked to another account.",
                    "record": {
                        "student_id": record.student_id,
                        "first_name": record.first_name,
                        "middle_name": record.middle_name,
                        "last_name": record.last_name,
                        "email": record.personal_email or "",
                        "degree": record.degree,
                        "year_graduated": record.year_graduated,
                        "college": record.college,
                    },
                }
            if not user.linked_student_id:
                user.linked_student_id = record.student_id
                user.is_verified = True
                db.commit()
            matched = True
            match_type = "email_name"
            reason = f"Matched by email/name • Student ID {record.student_id}"
    return {
        "matched": matched,
        "match_type": match_type,
        "reason": reason,
        "record": {
            "student_id": record.student_id,
            "first_name": record.first_name,
            "middle_name": record.middle_name,
            "last_name": record.last_name,
            "email": record.personal_email or "",
            "degree": record.degree,
            "year_graduated": record.year_graduated,
            "college": record.college,
        }
        if record
        else None,
    }


@router.post("/approvals/{account_id}/approve")
def approve(account_id: int, admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    user = db.get(Account, account_id)
    if not user or user.role != "Alumni":
        raise HTTPException(status_code=404, detail="Alumni account not found.")
    if user.status == "Active":
        if not user.is_verified:
            user.is_verified = True
            db.commit()
        return {"ok": True, "status": user.status, "already_processed": True}
    if user.status != "Pending":
        raise HTTPException(
            status_code=409,
            detail="This registration is no longer pending and cannot be approved.",
        )
    old = user.status
    user.status = "Active"
    user.is_verified = True
    _log(db, admin.id, "APPROVE", account_id, {"status": old}, {"status": "Active"})
    recipient = user.personal_email
    recipient_name = _account_name(user)
    db.commit()
    send_account_status_email(to_address=recipient, name=recipient_name, status="Active")
    return {"ok": True, "status": user.status}


@router.post("/approvals/{account_id}/reject")
def reject(account_id: int, payload: RejectRequest, admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    user = db.get(Account, account_id)
    if not user or user.role != "Alumni":
        raise HTTPException(status_code=404, detail="Alumni account not found.")
    if user.status == "Rejected":
        return {"ok": True, "status": user.status, "already_processed": True}
    if user.status != "Pending":
        raise HTTPException(
            status_code=409,
            detail="This registration is no longer pending and cannot be rejected.",
        )
    old = user.status
    user.status = "Rejected"
    user.rejection_reason = payload.reason
    _log(db, admin.id, "REJECT", account_id, {"status": old}, {"status": "Rejected", "reason": payload.reason})
    recipient = user.personal_email
    recipient_name = _account_name(user)
    reason = payload.reason
    db.commit()
    send_account_status_email(
        to_address=recipient,
        name=recipient_name,
        status="Rejected",
        reason=reason,
    )
    return {"ok": True, "status": user.status}


@router.get("/certifications/{certification_id}/file")
def admin_certification_file(
    certification_id: int,
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    row = certification_file(db, certification_id)
    path = resolve_stored_file(row.certificate_path, missing_detail="Certificate file is no longer on disk.")
    return FileResponse(
        path,
        media_type=row.certificate_mime or "application/octet-stream",
        filename=safe_download_name(row.certificate_filename or "certificate"),
        content_disposition_type="inline",
    )


@router.get("/resumes/{resume_id}/file")
def download_resume(resume_id: int, admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    row = db.get(Resume, resume_id)
    if not row:
        raise HTTPException(status_code=404, detail="Resume not found.")
    path = resolve_stored_file(row.stored_path, missing_detail="Resume file is no longer on disk.")
    return FileResponse(path, filename=safe_download_name(row.original_filename))


@router.get("/academic-filters")
def academic_filters(admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    ensure_academic_catalog(db, commit=True)
    return academic_catalog(db)


@router.get("/tracer")
def tracer_records(
    email: str = "",
    alignment: str = "All",
    college: str = "",
    course: str = "",
    submitted_on: str = "",
    page: int = 1,
    page_size: int = 20,
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    page = max(1, page)
    page_size = min(50, max(5, page_size))
    q = (
        db.query(TracerSubmission)
        .join(Account, Account.id == TracerSubmission.account_id)
        .options(
            joinedload(TracerSubmission.account).joinedload(Account.profile),
            joinedload(TracerSubmission.account).joinedload(Account.university_record),
        )
    )
    terms = [part for part in email.split() if part.strip()]
    if terms:
        saved_first = cast(func.coalesce(func.json_extract(TracerSubmission.data_json, "$.first_name"), ""), String)
        saved_middle = cast(func.coalesce(func.json_extract(TracerSubmission.data_json, "$.middle_name"), ""), String)
        saved_last = cast(func.coalesce(func.json_extract(TracerSubmission.data_json, "$.last_name"), ""), String)
        space = literal(" ")
        saved_name = func.trim(saved_first.op("||")(space).op("||")(func.trim(saved_middle.op("||")(space).op("||")(saved_last))))
        account_name = func.trim(
            cast(func.coalesce(Account.first_name, ""), String).op("||")(space).op("||")(cast(func.coalesce(Account.last_name, ""), String))
        )
        for term in terms:
            like = f"%{term}%"
            profile_match = (
                db.query(AlumniProfile.id)
                .filter(
                    AlumniProfile.account_id == Account.id,
                    or_(
                        AlumniProfile.first_name.ilike(like),
                        AlumniProfile.middle_name.ilike(like),
                        AlumniProfile.last_name.ilike(like),
                    ),
                )
                .exists()
            )
            registry_match = (
                db.query(UniversityRecord.student_id)
                .filter(
                    UniversityRecord.student_id == Account.linked_student_id,
                    or_(
                        UniversityRecord.first_name.ilike(like),
                        UniversityRecord.middle_name.ilike(like),
                        UniversityRecord.last_name.ilike(like),
                    ),
                )
                .exists()
            )
            q = q.filter(
                or_(
                    Account.personal_email.ilike(like),
                    Account.first_name.ilike(like),
                    Account.last_name.ilike(like),
                    account_name.ilike(like),
                    TracerSubmission.student_id.ilike(like),
                    Account.linked_student_id.ilike(like),
                    saved_first.ilike(like),
                    saved_middle.ilike(like),
                    saved_last.ilike(like),
                    saved_name.ilike(like),
                    profile_match,
                    registry_match,
                )
            )
    if alignment and alignment != "All":
        q = q.filter(TracerSubmission.alignment_status == alignment)
    day = (submitted_on or "").strip()
    if day:
        try:
            datetime.strptime(day, "%Y-%m-%d")
        except ValueError as exc:
            raise HTTPException(status_code=400, detail="Use a YYYY-MM-DD submission date.") from exc
        q = q.filter(cast(func.date(TracerSubmission.submitted_at), String) == day)
    q = apply_tracer_academic_filters(q, college, course)
    total = q.count()
    rows = q.order_by(TracerSubmission.submitted_at.desc()).offset((page - 1) * page_size).limit(page_size).all()
    items = []
    for row in rows:
        items.append(_tracer_payload(row))
    return {"items": items, "total": total, "page": page, "page_size": page_size}


def _join_name(*parts: object) -> str:
    return " ".join(str(part).strip() for part in parts if str(part or "").strip())


def _saved_name(row: TracerSubmission, data: dict) -> str:
    account = row.account
    profile = account.profile if account else None
    if profile:
        name = _join_name(profile.first_name, profile.middle_name, profile.last_name)
        if name:
            return name
    registry = account.university_record if account else None
    if registry:
        name = _join_name(registry.first_name, registry.middle_name, registry.last_name)
        if name:
            return name
    if account:
        name = _join_name(account.first_name, account.last_name)
        if name:
            return name
    return _join_name(data.get("first_name"), data.get("middle_name"), data.get("last_name"))


def _saved_student_number(row: TracerSubmission, data: dict) -> str:
    account = row.account
    candidates = [
        row.student_id,
        account.linked_student_id if account else None,
        data.get("student_id"),
        data.get("student_number"),
    ]
    for value in candidates:
        text = str(value or "").strip()
        if text:
            return text
    return ""


def _tracer_payload(row: TracerSubmission) -> dict:
    data = row.data_json or {}
    current = resolve_current_employment(data)
    return {
        "id": row.id,
        "account_id": row.account_id,
        "name": _saved_name(row, data),
        "email": row.account.personal_email if row.account else "",
        "student_id": _saved_student_number(row, data),
        "job_title": current.occupation,
        "employer": current.employer,
        "employment_status": data.get("is_currently_employed") or data.get("first_stat") or "",
        "alignment": row.alignment_status,
        "score": row.alignment_score,
        "submitted_at": row.submitted_at,
        "soc_code": row.soc_code or data.get("soc_code") or "",
        "soc_description": data.get("soc_description") or "",
        "soc_category": data.get("soc_category") or "",
        "data": data,
        "extra_answers": row.extra_answers,
        "survey_version": row.survey_version or data.get("_survey_version"),
    }


@router.get("/tracer/{submission_id}")
def tracer_record(submission_id: int, admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    row = (
        db.query(TracerSubmission)
        .options(joinedload(TracerSubmission.account))
        .filter(TracerSubmission.id == submission_id)
        .first()
    )
    if not row:
        raise HTTPException(status_code=404, detail="Tracer record not found.")
    payload = _tracer_payload(row)
    payload["further_studies"] = serialize_studies(db, row.account_id)
    payload["profile_activity"] = serialize_activity(db, row.account_id)
    return payload


@router.post("/tracer/{submission_id}")
def update_tracer_record(
    submission_id: int,
    payload: GtsPayload,
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    row = (
        db.query(TracerSubmission)
        .options(joinedload(TracerSubmission.account))
        .filter(TracerSubmission.id == submission_id)
        .first()
    )
    if not row or not row.account:
        raise HTTPException(status_code=404, detail="Tracer record not found.")
    data = payload.model_dump()
    extra = data.pop("extra_answers", {}) or {}
    schema = get_published_schema(db)
    validate_required_answers(schema, data, extra)
    submission = persist_tracer_submission(db, row.account, data, extra_answers=extra, resume_id=data.get("resume_id"))
    db.commit()
    db.refresh(submission)
    updated = _tracer_payload(submission)
    updated["further_studies"] = serialize_studies(db, submission.account_id)
    updated["profile_activity"] = serialize_activity(db, submission.account_id)
    return updated


@router.get("/questions")
def list_questions(admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    survey = ensure_gts_survey(db)
    db.commit()
    draft = survey.draft_schema or {}
    extras = extras_payload(draft)
    core_sections = []
    items = []
    for section, _sub, question in iter_questions(draft):
        is_extra = question.get("storage") == "extra"
        extra_id = question.get("extra_key") if is_extra else question.get("id")
        items.append(
            {
                "id": extra_id,
                "schema_id": question.get("id"),
                "section_key": section.get("key"),
                "label": question.get("label"),
                "input_type": question.get("type"),
                "options": question.get("options") or [],
                "required": bool(question.get("required")),
                "order_index": 0,
                "is_core": not is_extra,
                "active": True,
                "system": bool(question.get("system")),
            }
        )
    for section in draft.get("sections") or []:
        labels = [
            question.get("label")
            for sub in section.get("subsections") or []
            for question in sub.get("questions") or []
            if question.get("storage") != "extra"
        ]
        core_sections.append(
            {
                "key": section.get("key"),
                "name": section.get("name"),
                "locked": False,
                "fields": labels,
            }
        )
    return {
        "core_sections": core_sections or CORE_GTS_SECTIONS,
        "question_types": QUESTION_TYPES,
        "items": extras,
        "all_items": items,
    }


@router.post("/questions")
def create_question(payload: QuestionIn, admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    if payload.input_type not in QUESTION_TYPES:
        raise HTTPException(status_code=400, detail="Unsupported question type.")
    if not payload.label.strip():
        raise HTTPException(status_code=400, detail="Question title is required.")
    return add_legacy_extra(db, admin, payload)


@router.patch("/questions/{question_id}")
def update_question(question_id: int, payload: QuestionUpdate, admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    survey = ensure_gts_survey(db)
    draft = clone_schema(survey.draft_schema)
    target = None
    for _section, _sub, question in iter_questions(draft):
        if str(question.get("extra_key") or "") == str(question_id) or question.get("id") == f"extra_{question_id}":
            target = question
            break
    if not target:
        raise HTTPException(status_code=404, detail="Question not found.")
    data = payload.model_dump(exclude_unset=True)
    if "label" in data and data["label"]:
        target["label"] = data["label"]
    if "input_type" in data and data["input_type"]:
        target["type"] = data["input_type"]
    if "options" in data and data["options"] is not None:
        target["options"] = data["options"]
    if "required" in data and data["required"] is not None:
        target["required"] = data["required"]
    if "section_key" in data and data["section_key"]:
        moved = target
        for section in draft.get("sections") or []:
            for sub in section.get("subsections") or []:
                sub["questions"] = [item for item in sub.get("questions") or [] if item.get("id") != moved.get("id")]
        dest = next((section for section in draft.get("sections") or [] if section.get("key") == data["section_key"]), None)
        if dest:
            dest.setdefault("subsections", [{"id": f"{dest['id']}_main", "name": "", "questions": []}])
            dest["subsections"][-1].setdefault("questions", []).append(moved)
    save_draft(db, admin, draft)
    _log(db, admin.id, "UPDATE_QUESTION", question_id, None, data)
    return {"ok": True}


@router.delete("/questions/{question_id}")
def delete_question(question_id: int, admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    delete_legacy_extra(db, admin, question_id)
    return {"ok": True}


@router.get("/soc")
def list_soc(
    q: str = "",
    page: int = 1,
    page_size: int = 25,
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    query = db.query(SocCode)
    if q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(
            or_(
                SocCode.description.ilike(like),
                SocCode.soc_code.ilike(like),
                SocCode.category.ilike(like),
                SocCode.title_patterns.ilike(like),
            )
        )
    page = max(1, page)
    page_size = min(100, max(5, page_size))
    total = query.count()
    rows = query.order_by(SocCode.soc_code).offset((page - 1) * page_size).limit(page_size).all()
    return {
        "items": [
            {
                "soc_code": r.soc_code,
                "code": r.code,
                "description": r.description,
                "category": r.category,
                "title_patterns": r.title_patterns,
                "degree_patterns": r.degree_patterns,
                "major_group": r.major_group,
            }
            for r in rows
        ],
        "total": total,
        "page": page,
        "page_size": page_size,
    }


@router.post("/soc")
def create_soc(payload: SocIn, admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    code = payload.soc_code.strip()
    if not code or not payload.description.strip():
        raise HTTPException(status_code=400, detail="SOC code and description are required.")
    if db.get(SocCode, code):
        raise HTTPException(status_code=400, detail="That SOC code already exists.")
    row = SocCode(
        soc_code=code,
        code=code,
        description=payload.description.strip(),
        category=payload.category,
        title_patterns=payload.title_patterns,
        degree_patterns=payload.degree_patterns,
        major_group=payload.major_group,
    )
    db.add(row)
    _log(db, admin.id, "CREATE_SOC", code, None, payload.model_dump())
    db.commit()
    invalidate_soc_cache()
    return {"ok": True, "soc_code": code, "item": _soc_item(row)}


_SOC_IMPORT_HEADERS = {
    "soc_code": "soc_code",
    "code": "soc_code",
    "psoc_code": "soc_code",
    "psoc_soc_code": "soc_code",
    "description": "description",
    "category": "category",
    "title_patterns": "title_patterns",
    "job_title_patterns": "title_patterns",
    "job_titles": "title_patterns",
    "degree_patterns": "degree_patterns",
    "degree_keywords": "degree_patterns",
    "major_group": "major_group",
    "psoc_major_group": "major_group",
}


def _soc_header_key(value) -> str:
    text = "_".join(str(value or "").strip().lower().replace("/", " ").replace("-", " ").split())
    return _SOC_IMPORT_HEADERS.get(text, "")


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


@router.post("/soc/import")
def import_soc(
    file: UploadFile = File(...),
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    name = (file.filename or "").lower()
    if not name.endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="Upload an Excel workbook (.xlsx).")
    data = file.file.read(5 * 1024 * 1024 + 1)
    if len(data) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="The workbook must be 5 MB or smaller.")
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="This file is not a valid Excel workbook.")
    try:
        sheet = book.active
        if sheet is None:
            raise HTTPException(status_code=400, detail="The workbook has no worksheet.")
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise HTTPException(status_code=400, detail="The workbook is empty.")
        columns = [_soc_header_key(cell) for cell in header]
        if "soc_code" not in columns or "description" not in columns:
            raise HTTPException(
                status_code=400,
                detail="The first row needs Code and Description columns.",
            )
        added = 0
        skipped = 0
        issues: list[str] = []
        seen: set[str] = set()
        for index, raw in enumerate(rows, start=2):
            if index > 2001:
                issues.append("Only the first 2,000 rows were read.")
                break
            values = [_cell_text(cell) for cell in raw]
            if not any(values):
                continue
            record = {key: "" for key in ("soc_code", "description", "category", "title_patterns", "degree_patterns", "major_group")}
            for key, value in zip(columns, values):
                if key:
                    record[key] = value
            code = record["soc_code"]
            description = record["description"]
            if not code or not description:
                issues.append(f"Row {index}: SOC code and description are required.")
                continue
            if len(code) > 20 or len(description) > 255 or len(record["category"]) > 120 or len(record["major_group"]) > 80:
                issues.append(f"Row {index}: A value is longer than the catalog allows.")
                continue
            if code in seen or db.get(SocCode, code):
                skipped += 1
                continue
            seen.add(code)
            db.add(
                SocCode(
                    soc_code=code,
                    code=code,
                    description=description,
                    category=record["category"],
                    title_patterns=record["title_patterns"][:4000],
                    degree_patterns=record["degree_patterns"][:4000],
                    major_group=record["major_group"],
                )
            )
            added += 1
    finally:
        book.close()
    if added:
        _log(db, admin.id, "IMPORT_SOC", file.filename or "soc.xlsx", None, {"added": added, "skipped": skipped})
        db.commit()
        invalidate_soc_cache()
    else:
        db.rollback()
    if added:
        message = f"Added {added} occupational code{'s' if added != 1 else ''}."
    elif skipped:
        message = "No new codes were added. Matching codes already exist."
    else:
        message = "No occupational codes were added."
    return {
        "ok": True,
        "message": message,
        "added": added,
        "skipped": skipped,
        "invalid": len(issues),
        "issues": issues[:20],
    }


def _soc_item(row: SocCode) -> dict:
    return {
        "soc_code": row.soc_code,
        "code": row.code,
        "description": row.description,
        "category": row.category or "",
        "title_patterns": row.title_patterns or "",
        "degree_patterns": row.degree_patterns or "",
        "major_group": row.major_group or "",
    }


def _mapping_item(row: JobTitleMapping) -> dict:
    return {
        "id": row.id,
        "raw_title": row.raw_title,
        "normalized_title": row.normalized_title or "",
        "soc_code": row.soc_code or "",
        "admin_notes": row.admin_notes or "",
        "approved": bool(row.approved),
    }


@router.put("/soc/{soc_code}")
def update_soc(soc_code: str, payload: SocIn, admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    row = db.get(SocCode, soc_code)
    if not row:
        raise HTTPException(status_code=404, detail="SOC code not found.")
    code = payload.soc_code.strip()
    description = payload.description.strip()
    if not code or not description:
        raise HTTPException(status_code=400, detail="SOC code and description are required.")
    fields = {
        "description": description,
        "category": (payload.category or "").strip(),
        "title_patterns": (payload.title_patterns or "").strip(),
        "degree_patterns": (payload.degree_patterns or "").strip(),
        "major_group": (payload.major_group or "").strip(),
    }
    if code != row.soc_code:
        if db.get(SocCode, code) or db.query(SocCode).filter(SocCode.code == code).first():
            raise HTTPException(status_code=400, detail="That SOC code already exists.")
        replacement = SocCode(soc_code=code, code=code, **fields)
        db.add(replacement)
        db.flush()
        db.query(JobTitleMapping).filter(JobTitleMapping.soc_code == row.soc_code).update(
            {JobTitleMapping.soc_code: code},
            synchronize_session=False,
        )
        db.query(TracerSubmission).filter(TracerSubmission.soc_code == row.soc_code).update(
            {TracerSubmission.soc_code: code},
            synchronize_session=False,
        )
        db.delete(row)
        saved = replacement
    else:
        row.code = code
        for key, value in fields.items():
            setattr(row, key, value)
        saved = row
    _log(db, admin.id, "UPDATE_SOC", code, {"soc_code": soc_code}, {"soc_code": code, **fields})
    db.commit()
    db.refresh(saved)
    invalidate_soc_cache()
    return {"ok": True, "message": "SOC mapping updated.", "item": _soc_item(saved)}


@router.get("/job-mappings")
def list_mappings(
    q: str = "",
    page: int = 1,
    page_size: int = 25,
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    query = db.query(JobTitleMapping)
    if q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(
            or_(
                JobTitleMapping.raw_title.ilike(like),
                JobTitleMapping.soc_code.ilike(like),
                JobTitleMapping.admin_notes.ilike(like),
            )
        )
    page = max(1, page)
    page_size = min(100, max(5, page_size))
    total = query.count()
    rows = query.order_by(JobTitleMapping.raw_title.asc()).offset((page - 1) * page_size).limit(page_size).all()
    return {
        "items": [_mapping_item(row) for row in rows],
        "total": total,
        "page": page,
        "page_size": page_size,
    }


@router.put("/job-mappings/{mapping_id}")
def update_mapping(
    mapping_id: int,
    payload: MappingIn,
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    row = db.get(JobTitleMapping, mapping_id)
    if not row:
        raise HTTPException(status_code=404, detail="Job title mapping not found.")
    title = payload.raw_title.strip()
    code = payload.soc_code.strip()
    if not title:
        raise HTTPException(status_code=400, detail="Job title is required.")
    if not db.get(SocCode, code):
        raise HTTPException(status_code=400, detail="Unknown SOC code.")
    duplicate = (
        db.query(JobTitleMapping)
        .filter(JobTitleMapping.raw_title == title, JobTitleMapping.id != row.id)
        .first()
    )
    if duplicate:
        raise HTTPException(status_code=400, detail="That job title is already mapped.")
    row.raw_title = title
    row.normalized_title = title.lower()
    row.soc_code = code
    row.admin_notes = payload.admin_notes or ""
    row.approved = True
    row.last_modified_by = admin.id
    _log(db, admin.id, "UPDATE_JOB_MAPPING", title, None, payload.model_dump())
    db.commit()
    db.refresh(row)
    return {"ok": True, "message": "Job title mapping updated.", "item": _mapping_item(row)}


@router.post("/job-mappings")
def create_mapping(payload: MappingIn, admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    title = payload.raw_title.strip()
    if not title:
        raise HTTPException(status_code=400, detail="Job title is required.")
    if not db.get(SocCode, payload.soc_code):
        raise HTTPException(status_code=400, detail="Unknown SOC code.")
    existing = db.query(JobTitleMapping).filter(JobTitleMapping.raw_title == title).first()
    if existing:
        existing.soc_code = payload.soc_code
        existing.normalized_title = title.lower()
        existing.admin_notes = payload.admin_notes
        existing.last_modified_by = admin.id
        existing.approved = True
    else:
        db.add(
            JobTitleMapping(
                raw_title=title,
                normalized_title=title.lower(),
                soc_code=payload.soc_code,
                approved=True,
                admin_notes=payload.admin_notes,
                last_modified_by=admin.id,
            )
        )
    _log(db, admin.id, "MAP_JOB_TITLE", title, None, payload.model_dump())
    db.commit()
    return {"ok": True}


@router.get("/reports")
def reports(admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    alumni = db.query(Account).filter(Account.role == "Alumni")
    by_status = {
        "Active": alumni.filter(Account.status == "Active").count(),
        "Pending": alumni.filter(Account.status == "Pending").count(),
        "Rejected": alumni.filter(Account.status == "Rejected").count(),
    }
    by_year = (
        db.query(AlumniProfile.year_graduated, func.count(AlumniProfile.id))
        .group_by(AlumniProfile.year_graduated)
        .order_by(AlumniProfile.year_graduated)
        .all()
    )
    latest_rows = latest_submissions(db)
    related = 0
    employed = 0
    for row in latest_rows:
        profile = row.account.profile if row.account else None
        data = apply_profile_unemployment(row.data_json or {}, profile)
        if data.get("is_currently_employed") == "Yes":
            employed += 1
        if _relatedness(data, row.alignment_status or "") == "related":
            related += 1
    return {
        "registered": alumni.count(),
        "approved": by_status["Active"],
        "pending": by_status["Pending"],
        "rejected": by_status["Rejected"],
        "by_status": by_status,
        "by_graduation_year": [{"year": y or "Unknown", "count": c} for y, c in by_year],
        "currently_employed_submissions": employed,
        "degree_related_submissions": related,
        "alumni_with_submissions": len(latest_rows),
        "total_submissions": db.query(TracerSubmission).count(),
        "options": filter_options(db),
    }


@router.get("/reports/employment")
def employment_report(
    college: str = "",
    program: str = "",
    course: str = "",
    year: str = "",
    location: str = "",
    employment: str = "",
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    try:
        return employment_analytics(
            db,
            college=college,
            program=program or course,
            year=year,
            location=location,
            employment=employment,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


def _oaaps_filters(
    college: str = "",
    program: str = "",
    year: str = "",
    period_from: str = "",
    period_to: str = "",
    award: str = "",
    further: str = "",
):
    try:
        submitted_from = parse_report_date(period_from)
        submitted_to = parse_report_date(period_to, end_of_day=True)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Reporting period dates must use YYYY-MM-DD.") from exc
    if submitted_from and submitted_to and submitted_from > submitted_to:
        raise HTTPException(status_code=400, detail="Reporting period start date must be on or before the end date.")
    return {
        "college": college,
        "program": program,
        "year": year,
        "period_from": period_from,
        "period_to": period_to,
        "award": award,
        "further": further,
        "submitted_from": submitted_from,
        "submitted_to": submitted_to,
    }


@router.get("/reports/oaaps")
def oaaps_reports(
    college: str = "",
    program: str = "",
    year: str = "",
    period_from: str = "",
    period_to: str = "",
    award: str = "",
    further: str = "",
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    return build_oaaps_report(db, **_oaaps_filters(college, program, year, period_from, period_to, award, further))


@router.get("/reports/oaaps/export")
def export_oaaps_report(
    college: str = "",
    program: str = "",
    year: str = "",
    period_from: str = "",
    period_to: str = "",
    award: str = "",
    further: str = "",
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    report = build_oaaps_report(db, **_oaaps_filters(college, program, year, period_from, period_to, award, further))
    payload = render_oaaps_docx(report)
    filename = f"AUF-OAAPS-Graduate-Productivity-Report-{datetime.now().strftime('%Y%m%d')}.docx"
    return StreamingResponse(
        iter([payload]),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/reports/export")
def export_reports(
    college: str = "",
    program: str = "",
    course: str = "",
    year: str = "",
    location: str = "",
    employment: str = "",
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(
        [
            "email",
            "student_id",
            "name",
            "degree",
            "year_graduated",
            "status",
            "job_title",
            "employer",
            "alignment",
            "submitted_at",
        ]
    )
    try:
        rows = []
        for row in _latest_tracer_rows(db):
            prepared = prepare_employment_row(row, college, program or course, year, location, employment)
            if prepared is None:
                continue
            rows.append((row, prepared[0]))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    for row, data in rows:
        writer.writerow(
            [
                csv_safe(row.account.personal_email if row.account else ""),
                csv_safe(row.student_id or ""),
                csv_safe(
                    " ".join(
                        p
                        for p in [
                            data.get("first_name"),
                            data.get("middle_name"),
                            data.get("last_name"),
                        ]
                        if p
                    )
                ),
                csv_safe(data.get("degree", "")),
                csv_safe(data.get("year_graduated", "")),
                csv_safe(row.account.status if row.account else ""),
                csv_safe(resolve_current_employment(data).occupation),
                csv_safe(resolve_current_employment(data).employer),
                csv_safe(row.alignment_status),
                csv_safe(row.submitted_at.isoformat() if row.submitted_at else ""),
            ]
        )
    output.seek(0)
    filename = f"careersense-tracer-{datetime.now().strftime('%Y%m%d')}.csv"
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/university-records")
def list_university_records(
    q: str = "",
    college: str = "",
    course: str = "",
    page: int = 1,
    page_size: int = 25,
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    query = db.query(UniversityRecord)
    if q.strip():
        like = f"%{q.strip()}%"
        query = query.filter(
            or_(
                UniversityRecord.student_id.ilike(like),
                UniversityRecord.personal_email.ilike(like),
                UniversityRecord.first_name.ilike(like),
                UniversityRecord.middle_name.ilike(like),
                UniversityRecord.last_name.ilike(like),
                UniversityRecord.college.ilike(like),
                UniversityRecord.degree.ilike(like),
                UniversityRecord.course_code.ilike(like),
            )
        )
    query = apply_registry_academic_filters(query, college, course)
    page = max(1, page)
    page_size = min(100, max(5, page_size))
    total = query.count()
    rows = query.order_by(
        UniversityRecord.last_name, UniversityRecord.first_name, UniversityRecord.student_id
    ).offset((page - 1) * page_size).limit(page_size).all()
    return {
        "items": [
            {
                "student_id": r.student_id,
                "first_name": r.first_name,
                "middle_name": r.middle_name,
                "last_name": r.last_name,
                "personal_email": r.personal_email or "",
                "degree": r.degree,
                "year_graduated": r.year_graduated,
                "course_code": r.course_code,
                "college": r.college,
            }
            for r in rows
        ],
        "total": total,
        "page": page,
        "page_size": page_size,
        "last_import": latest_registry_import(db),
    }


@router.post("/university-records")
def create_university_record(payload: UniversityRecordIn, admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    if db.get(UniversityRecord, payload.student_id.strip()):
        raise HTTPException(status_code=400, detail="Student ID already exists.")
    email = str(payload.personal_email).strip().lower()
    if db.query(UniversityRecord).filter(UniversityRecord.personal_email == email).first():
        raise HTTPException(status_code=400, detail="A graduate record already uses this email.")
    try:
        college, degree, course_code = resolve_program_selection(
            db, payload.college, payload.degree, payload.course_code
        )
    except ProgramSelectionError as exc:
        raise HTTPException(status_code=400, detail=exc.message) from exc
    known = any(
        item["college"].casefold() == college.casefold() and item["course"].casefold() == degree.casefold()
        for item in academic_programs(db)
    )
    owner = None if known else course_code_owner(db, course_code)
    if owner:
        raise HTTPException(
            status_code=400,
            detail=f"Course code {course_code} is already used by {owner['course']}.",
        )
    row = UniversityRecord(
        student_id=payload.student_id.strip(),
        first_name=payload.first_name.strip(),
        middle_name=payload.middle_name.strip(),
        last_name=payload.last_name.strip(),
        personal_email=email,
        degree=degree,
        year_graduated=payload.year_graduated.strip(),
        course_code=course_code,
        college=college,
    )
    db.add(row)
    db.flush()
    sync_catalog_from_records(db)
    _log(db, admin.id, "CREATE_UNIVERSITY_RECORD", row.student_id, None, payload.model_dump())
    db.commit()
    return {"ok": True, "student_id": row.student_id, "course_code": row.course_code, "college": row.college, "degree": row.degree}


def _read_upload(file: UploadFile) -> bytes:
    return file.file.read()


@router.post("/university-records/import/preview")
def preview_university_import(
    file: UploadFile = File(...),
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    try:
        return preview_registry_import(db, _read_upload(file), file.filename or "")
    except RegistryImportError as exc:
        raise HTTPException(status_code=exc.status_code, detail=exc.message) from exc


@router.post("/university-records/import")
def commit_university_import(
    file: UploadFile = File(...),
    admin: Account = Depends(require_admin),
    db: Session = Depends(get_db),
):
    try:
        result = commit_registry_import(db, _read_upload(file), file.filename or "", admin.id)
    except RegistryImportError as exc:
        raise HTTPException(status_code=exc.status_code, detail=exc.message) from exc
    if result["summary"]["imported"]:
        _log(db, admin.id, "IMPORT_UNIVERSITY_RECORDS", result.get("filename") or "", None, result["summary"])
        db.commit()
    return result


@router.get("/logs")
def logs(admin: Account = Depends(require_admin), db: Session = Depends(get_db)):
    rows = db.query(AdminLog).order_by(AdminLog.created_at.desc()).limit(100).all()
    return [
        {
            "id": r.id,
            "admin_id": r.admin_id,
            "action_type": r.action_type,
            "target_id": r.target_id,
            "new_value": r.new_value,
            "created_at": r.created_at,
        }
        for r in rows
    ]
