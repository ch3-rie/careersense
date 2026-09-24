from io import BytesIO
from typing import Any, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

NAVY = RGBColor(0x0B, 0x2E, 0x59)
GOLD = RGBColor(0xB5, 0x8A, 0x3A)


def _set_run(run, *, size=11, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _shade(cell, hex_color: str):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _cell_text(cell, text, *, bold=False, size=10, align="left"):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = {
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    run = paragraph.add_run(str(text))
    _set_run(run, size=size, bold=bold)


def _add_table(doc: Document, headers: list[str], rows: list[list[Any]], *, numeric_from: int = 1):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        _cell_text(cell, header, bold=True, size=10, align="center")
        _shade(cell, "0B2E59")
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for r_index, row in enumerate(rows, start=1):
        is_total = str(row[0]).upper() == "TOTAL"
        for c_index, value in enumerate(row):
            align = "left" if c_index < numeric_from else "right"
            cell = table.rows[r_index].cells[c_index]
            _cell_text(cell, value, bold=is_total, size=10, align=align)
            if is_total:
                _shade(cell, "F4E6C4")
    doc.add_paragraph()
    return table


def _fmt(metric: Optional[dict], kind: str = "count") -> str:
    if metric is None:
        return "—"
    if isinstance(metric, dict) and metric.get("available") is False:
        return "Not Available" if kind == "pct" else "—"
    if kind == "pct":
        value = metric.get("percentage") if isinstance(metric, dict) else metric
        if value is None:
            return "Not Available"
        return f"{float(value):.2f}%"
    if isinstance(metric, dict):
        count = metric.get("count")
        return "—" if count is None else str(count)
    return str(metric)


def _heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    _set_run(run, size=13 if level == 1 else 12, bold=True, color=NAVY)


def _body(doc: Document, text: str, *, italic=False, bold=False, size=11):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    _set_run(run, size=size, bold=bold)
    run.italic = italic


def render_oaaps_docx(report: dict[str, Any]) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run(report["title"])
    _set_run(run, size=18, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gold = subtitle.add_run("Angeles University Foundation  ·  Office of Alumni Affairs and Placement Services")
    _set_run(gold, size=10, color=GOLD)

    header = report["header"]
    _add_table(
        doc,
        ["Item", "Value"],
        [
            ["Reporting Period", header["reporting_period"]],
            ["Graduating Batch/Cohort", header["batch_cohort"]],
            ["College/Program", header["college_program"]],
            ["Total Graduates", header["total_graduates"]],
            ["Number of Graduates Traced", header["graduates_traced"]],
            ["Tracer Study Response Rate", _fmt(header["response_rate"], "pct")],
        ],
        numeric_from=1,
    )

    _heading(doc, "I. OVERALL GRADUATE PRODUCTIVITY")
    _heading(doc, "A. Executive Summary", 2)
    _add_table(
        doc,
        ["Key Indicator", "Total", "Percentage"],
        [[row["indicator"], _fmt(row), _fmt(row, "pct")] for row in report["executive_summary"]],
    )
    _body(doc, f"Overall Graduate Productivity Rate: {report['overall_productivity_rate']:.2f}%", bold=True)
    _body(doc, report["formula"], italic=True, size=10)

    profile = report["profile"]
    _heading(doc, "B. Graduate Productivity Profile", 2)
    _body(
        doc,
        (
            f"Total Graduates: {profile['total_graduates']}     "
            f"Traced Graduates: {profile['traced']}     "
            f"Productive Graduates: {profile['productive']}     "
            f"Not Yet Productively Engaged: {profile['not_yet_engaged']}"
        ),
    )
    _body(doc, "Productivity Distribution:", bold=True)
    dist = report["distribution"]
    _body(doc, f"Workforce Integration: {dist['workforce']:.2f}%")
    _body(doc, f"Professional Advancement: {dist['professional']:.2f}%")
    _body(doc, f"Academic Advancement: {dist['academic']:.2f}%")

    workforce = report["workforce"]
    _heading(doc, "II. WORKFORCE INTEGRATION")
    _body(doc, workforce["description"], italic=True)
    _heading(doc, "A. Employment Status", 2)
    _add_table(
        doc,
        ["Employment Indicator", "No. of Graduates", "%"],
        [[row["indicator"], _fmt(row), _fmt(row, "pct")] for row in workforce["employment_status"]],
    )
    _heading(doc, "B. Employment Relevance", 2)
    _add_table(
        doc,
        ["Indicator", "No.", "%"],
        [[row["indicator"], _fmt(row), _fmt(row, "pct")] for row in workforce["employment_relevance"]],
    )
    _heading(doc, "C. Employment Quality", 2)
    _add_table(
        doc,
        ["Indicator", "No.", "%"],
        [[row["indicator"], _fmt(row), _fmt(row, "pct")] for row in workforce["employment_quality"]],
    )
    _heading(doc, "D. Time-to-Employment", 2)
    _add_table(
        doc,
        ["Indicator", "No.", "%"],
        [[row["indicator"], _fmt(row), _fmt(row, "pct")] for row in workforce["time_to_employment"]],
    )
    _body(doc, f"Key Workforce Integration Rate: {workforce['key_rate']:.2f}%", bold=True)

    professional = report["professional"]
    _heading(doc, "III. PROFESSIONAL ADVANCEMENT")
    _body(doc, professional["description"], italic=True)
    _heading(doc, "A. Career Progression", 2)
    _add_table(
        doc,
        ["Indicator", "No. of Graduates", "%"],
        [[row["indicator"], _fmt(row), _fmt(row, "pct")] for row in professional["career_progression"]],
    )
    _heading(doc, "B. Professional Credentials and Development", 2)
    _add_table(
        doc,
        ["Indicator", "No.", "%"],
        [[row["indicator"], _fmt(row), _fmt(row, "pct")] for row in professional["credentials"]],
    )
    _heading(doc, "C. Professional Recognition", 2)
    _body(doc, "Top Graduate Achievements:")
    _add_table(
        doc,
        ["Graduate", "Program/Batch", "Organization", "Position/Achievement", "Year"],
        [["Not Available — CareerSense does not currently collect this information.", "", "", "", ""]],
        numeric_from=5,
    )
    _body(doc, f"Professional Advancement Rate: {professional['key_rate']:.2f}%", bold=True)

    academic = report["academic"]
    _heading(doc, "IV. ACADEMIC ADVANCEMENT")
    _body(doc, academic["description"], italic=True)
    _heading(doc, "A. Further Education", 2)
    _add_table(
        doc,
        ["Academic Outcome", "No. of Graduates", "%"],
        [[row["indicator"], _fmt(row), _fmt(row, "pct")] for row in academic["further_education"]],
    )
    _heading(doc, "B. Academic and Research Achievements", 2)
    _add_table(
        doc,
        ["Indicator", "No.", "%"],
        [[row["indicator"], _fmt(row), _fmt(row, "pct")] for row in academic["achievements"]],
    )
    _body(doc, f"Academic Advancement Rate: {academic['key_rate']:.2f}%", bold=True)

    _heading(doc, "V. OVERALL GRADUATE PRODUCTIVITY SCORECARD")
    _add_table(
        doc,
        ["Dimension", "Key Indicator", "Target", "Actual Rate", "Status"],
        [
            [
                row["dimension"],
                row["key_indicator"],
                "Not Set" if row["target"] is None else row["target"],
                f"{row['actual_rate']:.2f}%",
                "Not Set" if row["status"] is None else row["status"],
            ]
            for row in report["scorecard"]
        ],
        numeric_from=2,
    )
    _body(doc, "Overall Productivity Formula", bold=True)
    _body(doc, report["formula"], italic=True)

    _heading(doc, "VI. GRADUATE PRODUCTIVITY BY COLLEGE/PROGRAM/YEAR")
    _add_table(
        doc,
        [
            "College/Program",
            "Batch Year",
            "Graduates",
            "Traced",
            "Workforce",
            "Professional",
            "Academic",
            "Overall Productivity",
        ],
        [
            [
                row["college_program"],
                row["batch_year"],
                row["graduates"],
                row["traced"],
                _fmt(row["workforce"], "pct"),
                _fmt(row["professional"], "pct"),
                _fmt(row["academic"], "pct"),
                _fmt(row["overall_productivity"], "pct"),
            ]
            for row in report["by_program"]
        ],
        numeric_from=1,
    )
    _body(
        doc,
        "Indicators marked Not Available or — are not collected in the current Graduate Tracer Survey and were not inferred.",
        italic=True,
        size=9,
    )

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
