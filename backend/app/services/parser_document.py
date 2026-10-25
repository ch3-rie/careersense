"""Layout-aware resume document extraction (PDF, DOCX, TXT) with optional OCR."""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from typing import Any, Protocol

logger = logging.getLogger(__name__)

MIN_TEXT_CHARS = 40
MIN_TEXT_WORDS = 5
DEFAULT_MAX_PDF_PAGES = 20
OCR_MIN_CONFIDENCE = 30
_DATE_ONLY = re.compile(
    r"(?:(?:19|20)\d{2}|present|current|now|ongoing|jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|"
    r"apr(?:il)?|may|jun(?:e)?|jul(?:y)?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|"
    r"nov(?:ember)?|dec(?:ember)?)",
    re.I,
)

try:
    import fitz
except ImportError:  # pragma: no cover
    fitz = None

try:
    from docx import Document
    from docx.oxml.ns import qn

    DOCX_AVAILABLE = True
except ImportError:  # pragma: no cover
    Document = None
    qn = None
    DOCX_AVAILABLE = False


@dataclass
class TextBlock:
    text: str
    page: int = 0
    x0: float = 0.0
    y0: float = 0.0
    x1: float = 0.0
    y1: float = 0.0
    font_size: float = 0.0
    kind: str = "text"

    @property
    def width(self) -> float:
        return max(0.0, self.x1 - self.x0)


@dataclass
class DocumentView:
    text: str = ""
    blocks: list[TextBlock] = field(default_factory=list)
    source: str = "txt"
    filename: str = ""
    page_count: int = 1
    columns: int = 1
    warnings: list[str] = field(default_factory=list)
    ocr_confidence: float | None = None
    ocr_pages: list[int] = field(default_factory=list)
    quality: dict = field(default_factory=dict)

    @classmethod
    def from_text(cls, text: str, filename: str = "resume.txt", source: str = "txt") -> "DocumentView":
        cleaned = clean_text(text)
        blocks = [
            TextBlock(text=line, page=0, y0=float(index), y1=float(index) + 1, kind="text")
            for index, line in enumerate(cleaned.splitlines())
            if line.strip()
        ]
        return cls(text=cleaned, blocks=blocks, source=source, filename=filename, page_count=1, columns=1)


class OcrEngine(Protocol):
    def image_to_text(self, image_bytes: bytes) -> str:
        ...


class NullOcrEngine:
    def image_to_text(self, image_bytes: bytes) -> str:
        return ""


def _configure_tesseract(pytesseract_module) -> None:
    """Use TESSERACT_CMD when set. Otherwise pytesseract searches PATH (Linux and Windows)."""
    import os

    command = (os.environ.get("TESSERACT_CMD") or "").strip()
    if command:
        pytesseract_module.pytesseract.tesseract_cmd = command


class TesseractOcrEngine:
    def image_to_text(self, image_bytes: bytes) -> str:
        text, _confidence = self.recognize(image_bytes)
        return text

    def recognize(self, image_bytes: bytes) -> tuple[str, float | None]:
        import pytesseract
        from PIL import Image

        _configure_tesseract(pytesseract)
        with Image.open(io.BytesIO(image_bytes)) as image:
            if image.mode not in {"RGB", "L"}:
                image = image.convert("RGB")
            try:
                data = pytesseract.image_to_data(image, output_type=pytesseract.Output.DICT)
                words = [word for word in data.get("text") or [] if str(word).strip()]
                confidences = []
                for raw in data.get("conf") or []:
                    try:
                        value = float(raw)
                    except (TypeError, ValueError):
                        continue
                    if value >= 0:
                        confidences.append(value)
                text = " ".join(words).strip() or (pytesseract.image_to_string(image) or "")
                confidence = round(sum(confidences) / len(confidences), 1) if confidences else None
                return text, confidence
            except Exception:
                logger.exception("Tesseract confidence data was unavailable; using plain OCR text.")
                return pytesseract.image_to_string(image) or "", None


_ocr_override: OcrEngine | None = None


def set_ocr_engine_for_tests(engine: OcrEngine | None) -> None:
    global _ocr_override
    _ocr_override = engine


def ocr_engine_available() -> bool:
    try:
        import pytesseract

        _configure_tesseract(pytesseract)
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def get_ocr_engine() -> OcrEngine:
    if _ocr_override is not None:
        return _ocr_override
    if ocr_engine_available():
        return TesseractOcrEngine()
    return NullOcrEngine()


def clean_text(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\x00", " ").replace("\u00a0", " ")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def assess_text_quality(text: str) -> dict:
    """Quality signals for selectable text. Sufficiency is not a count of characters alone."""
    cleaned = clean_text(text)
    compact = re.sub(r"\s+", "", cleaned)
    letters = re.sub(r"[^A-Za-z0-9]+", "", cleaned)
    words = re.findall(r"[A-Za-z]{2,}", cleaned)
    replacement = cleaned.count("\ufffd")
    controls = len(re.findall(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", text or ""))
    noise = len(re.findall(r"[^\w\s.,;:'\"()/\-+&@#%*•·|–—₱$]", cleaned))
    alnum_ratio = (len(letters) / len(compact)) if compact else 0.0
    noise_ratio = (noise / len(compact)) if compact else 0.0
    most = max((compact.count(char) for char in set(compact)), default=0)
    repeat_ratio = (most / len(compact)) if compact else 1.0
    base = len(letters) >= MIN_TEXT_CHARS and len(words) >= MIN_TEXT_WORDS
    sufficient = (
        base
        and alnum_ratio >= 0.35
        and noise_ratio <= 0.30
        and repeat_ratio < 0.50
        and replacement + controls <= 4
    )
    return {
        "characters": len(letters),
        "words": len(words),
        "alnum_ratio": round(alnum_ratio, 3),
        "noise_ratio": round(noise_ratio, 3),
        "repeat_ratio": round(repeat_ratio, 3),
        "sufficient": sufficient,
    }


def text_is_sufficient(text: str) -> bool:
    return bool(assess_text_quality(text)["sufficient"])


def max_pdf_pages() -> int:
    try:
        from app.config import get_settings

        return max(1, int(get_settings().parser_max_pdf_pages))
    except Exception:
        return DEFAULT_MAX_PDF_PAGES


def extract_document(filename: str, file_bytes: bytes, *, enable_ocr: bool = True) -> DocumentView:
    name = (filename or "resume.txt").lower()
    if not file_bytes:
        return DocumentView(source="empty", filename=filename, page_count=0)
    try:
        if name.endswith(".pdf"):
            view = _extract_pdf(filename, file_bytes)
            if view.source == "failed":
                return _with_quality(view)
            if enable_ocr:
                view = _fill_sparse_pages(filename, file_bytes, view)
            return _with_quality(view)
        if name.endswith(".docx"):
            return _with_quality(_extract_docx(filename, file_bytes))
        if name.endswith(".txt") or name.endswith(".md"):
            text = file_bytes.decode("utf-8", errors="ignore")
            if not text.strip():
                text = file_bytes.decode("latin-1", errors="ignore")
            view = DocumentView.from_text(text, filename=filename)
            if not text_is_sufficient(view.text):
                view.source = "empty" if not view.text else view.source
            return _with_quality(view)
        view = DocumentView.from_text(file_bytes.decode("utf-8", errors="ignore"), filename=filename)
        view.source = "empty" if not view.text else view.source
        return _with_quality(view)
    except Exception:
        logger.exception("Document reading failed for %s", filename or "resume")
        return DocumentView(source="failed", filename=filename, warnings=["extraction_failed"])


def _with_quality(view: DocumentView) -> DocumentView:
    view.quality = assess_text_quality(view.text)
    return view


def _extract_pdf(filename: str, file_bytes: bytes) -> DocumentView:
    if fitz is None:
        return DocumentView(source="failed", filename=filename, warnings=["pymupdf_missing"])
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception:
        logger.exception("PDF could not be opened: %s", filename or "resume")
        return DocumentView(source="failed", filename=filename, warnings=["pdf_unreadable"])
    blocks: list[TextBlock] = []
    max_columns = 1
    warnings: list[str] = []
    try:
        if getattr(doc, "needs_pass", False):
            return DocumentView(source="failed", filename=filename, warnings=["encrypted_pdf"], page_count=doc.page_count)
        limit = max_pdf_pages()
        page_count = doc.page_count
        if page_count > limit:
            warnings.append("page_limit")
        for page_index, page in enumerate(doc):
            if page_index >= limit:
                break
            page_blocks, columns = _pdf_page_blocks(page, page_index)
            blocks.extend(page_blocks)
            max_columns = max(max_columns, columns)
        ordered_text = "\n".join(block.text for block in blocks if block.text.strip())
        source = "digital_pdf" if text_is_sufficient(ordered_text) else "empty"
        return DocumentView(
            text=clean_text(ordered_text),
            blocks=blocks,
            source=source,
            filename=filename,
            page_count=page_count,
            columns=max_columns,
            warnings=warnings,
        )
    finally:
        doc.close()


def _pdf_page_blocks(page: Any, page_index: int) -> tuple[list[TextBlock], int]:
    width = float(page.rect.width or 1)
    height = float(page.rect.height or 1)
    raw_blocks: list[TextBlock] = []
    payload = page.get_text("dict") or {}
    for block in payload.get("blocks", []):
        if block.get("type") != 0:
            continue
        bbox = block.get("bbox") or [0, 0, 0, 0]
        for line in block.get("lines", []):
            spans = line.get("spans") or []
            piece = "".join(span.get("text", "") for span in spans).strip()
            if not piece:
                continue
            line_box = line.get("bbox") or bbox
            max_size = max((float(span.get("size") or 0) for span in spans), default=0.0)
            raw_blocks.append(
                TextBlock(
                    text=piece,
                    page=page_index,
                    x0=float(line_box[0]),
                    y0=float(line_box[1]),
                    x1=float(line_box[2]),
                    y1=float(line_box[3]),
                    font_size=max_size,
                    kind="text",
                )
            )
    table_blocks, table_boxes = _pdf_table_blocks(page, page_index)
    if table_boxes:
        raw_blocks = [block for block in raw_blocks if not _block_inside(block, table_boxes)]
    raw_blocks.extend(table_blocks)
    ordered, columns = _order_blocks(raw_blocks, width, height)
    return ordered, columns


def _block_inside(block: TextBlock, boxes: list[tuple[float, float, float, float]]) -> bool:
    cx = (block.x0 + block.x1) / 2
    cy = (block.y0 + block.y1) / 2
    return any(x0 - 1 <= cx <= x1 + 1 and y0 - 1 <= cy <= y1 + 1 for x0, y0, x1, y1 in boxes)


def _pdf_table_blocks(page: Any, page_index: int) -> tuple[list[TextBlock], list[tuple[float, float, float, float]]]:
    blocks: list[TextBlock] = []
    boxes: list[tuple[float, float, float, float]] = []
    try:
        finder = page.find_tables()
        tables = finder.tables if finder is not None else []
    except Exception:
        return blocks, boxes
    for table in tables or []:
        try:
            bbox = getattr(table, "bbox", None) or [0, 0, 0, 0]
            rows = table.extract() or []
        except Exception:
            continue
        produced = False
        for row_index, row in enumerate(rows):
            cells = [str(cell).strip() for cell in (row or []) if str(cell or "").strip()]
            if not cells:
                continue
            y0 = float(bbox[1]) + row_index * 12
            produced = True
            blocks.append(
                TextBlock(
                    text=" | ".join(cells),
                    page=page_index,
                    x0=float(bbox[0]),
                    y0=y0,
                    x1=float(bbox[2]),
                    y1=y0 + 12,
                    kind="table_cell",
                )
            )
        if produced and len(bbox) == 4:
            boxes.append(tuple(float(value) for value in bbox))
    return blocks, boxes


def _is_rail_label(text: str) -> bool:
    """A short section heading, using the same heading list as field extraction."""
    words = re.findall(r"[A-Za-z]+", text or "")
    if not 1 <= len(words) <= 6:
        return False
    from app.services.parser_extract import _heading_key

    return _heading_key(text) is not None


def _is_heading_rail(column: list[TextBlock]) -> bool:
    """A narrow column of section labels, such as PROFILE / EXPERIENCE down the left side."""
    if len(column) < 3:
        return False
    labels = sum(1 for block in column if _is_rail_label(block.text))
    return labels / len(column) >= 0.6


def _merge_heading_rail(rail: list[TextBlock], body: list[TextBlock]) -> list[TextBlock]:
    """Read a heading beside the lines at the same height, instead of the whole rail first."""
    combined = list(rail) + list(body)
    # A side label often sits a few points below the first line of the section it names.
    return sorted(
        combined,
        key=lambda block: (
            block.y0 - (10 if _is_rail_label(block.text) else 0),
            0 if _is_rail_label(block.text) else 1,
            block.x0,
        ),
    )


def _order_blocks(blocks: list[TextBlock], page_width: float, page_height: float) -> tuple[list[TextBlock], int]:
    if not blocks:
        return [], 1
    header_band = page_height * 0.12
    full_width = []
    column_candidates = []
    for block in blocks:
        # The name, title, and contact lines sit in the top band even when they are not full width.
        if block.y0 <= header_band or _is_full_width(block, page_width):
            full_width.append(block)
        else:
            column_candidates.append(block)
    header = [block for block in full_width if block.y0 <= page_height * 0.24]
    footer = [block for block in full_width if block not in header]
    columns = _attach_date_columns(_cluster_columns(column_candidates, page_width))
    rails = [column for column in columns if _is_heading_rail(column)]
    ordered = sorted(header, key=lambda block: (block.y0, block.x0))
    if len(rails) == 1 and len(columns) > 1:
        rail = rails[0]
        others = [column for column in columns if column is not rail]
        body = [block for column in others for block in column]
        ordered.extend(_merge_heading_rail(rail, body))
        column_count = max(1, len(others))
    else:
        for column in columns:
            ordered.extend(sorted(column, key=lambda block: (block.y0, block.x0)))
        column_count = max(1, len(columns))
    ordered.extend(sorted(footer, key=lambda block: (block.y0, block.x0)))
    return ordered, column_count


def _block_is_date(block: TextBlock) -> bool:
    text = (block.text or "").strip()
    if not text or len(text) > 48:
        return False
    if not _DATE_ONLY.search(text):
        return False
    leftover = _DATE_ONLY.sub("", text)
    leftover = re.sub(r"[\s\-\u2013\u2014/|to.,]+", "", leftover, flags=re.I)
    return len(leftover) <= 2


def _column_is_dates(column: list[TextBlock]) -> bool:
    if len(column) < 2:
        return False
    dated = sum(1 for block in column if _block_is_date(block))
    return dated >= 2 and dated / len(column) >= 0.6


def _attach_date_columns(columns: list[list[TextBlock]]) -> list[list[TextBlock]]:
    """Keep a date rail with the column beside it instead of reading it as a later section."""
    kept: list[list[TextBlock]] = []
    for column in columns:
        if kept and _column_is_dates(column):
            target = kept[-1]
            for date_block in column:
                nearest = None
                best = 10**9
                for block in target:
                    if block.page != date_block.page:
                        continue
                    gap = abs(block.y0 - date_block.y0)
                    if gap < best:
                        best = gap
                        nearest = block
                if nearest is not None and best <= 40:
                    nearest.text = f"{nearest.text} | {date_block.text}".strip(" |")
                else:
                    target.append(date_block)
            continue
        kept.append(column)
    return kept


def _is_full_width(block: TextBlock, page_width: float) -> bool:
    if page_width <= 0:
        return True
    span = block.width / page_width
    starts_left = block.x0 <= page_width * 0.12
    ends_right = block.x1 >= page_width * 0.82
    # A long line that begins in a side column is still that column, not a page banner.
    return starts_left and ends_right and span >= 0.62


def _cluster_columns(blocks: list[TextBlock], page_width: float) -> list[list[TextBlock]]:
    if len(blocks) < 3 or page_width <= 0:
        return [blocks] if blocks else []
    xs = sorted(block.x0 for block in blocks)
    gaps = [(xs[index + 1] - xs[index], index) for index in range(len(xs) - 1)]
    # A single page number far to the right must not cancel the real column gap.
    for max_gap, gap_at in sorted(gaps, reverse=True):
        if max_gap < page_width * 0.16:
            break
        split_x = (xs[gap_at] + xs[gap_at + 1]) / 2
        left = [block for block in blocks if block.x0 < split_x]
        right = [block for block in blocks if block.x0 >= split_x]
        if len(left) < 2 or len(right) < 2:
            continue
        grouped = []
        for side in (left, right):
            nested = _cluster_columns(side, page_width * 0.7) if len(side) >= 6 else [side]
            grouped.extend(nested)
        return grouped or [blocks]
    return [blocks]


def _recognize_image(engine: OcrEngine, image_bytes: bytes) -> tuple[str, float | None]:
    recognize = getattr(engine, "recognize", None)
    if callable(recognize):
        try:
            text, confidence = recognize(image_bytes)
            return (text or "").strip(), confidence
        except Exception:
            logger.exception("OCR recognition failed")
            return "", None
    try:
        return (engine.image_to_text(image_bytes) or "").strip(), None
    except Exception:
        logger.exception("OCR recognition failed")
        return "", None


def _fill_sparse_pages(filename: str, file_bytes: bytes, view: DocumentView) -> DocumentView:
    """OCR pages that have little selectable text, including mixed digital and scanned PDFs."""
    if fitz is None:
        return view
    limit = min(view.page_count or 0, max_pdf_pages())
    sparse = [index for index in range(limit) if not text_is_sufficient(_page_text(view, index))]
    if not sparse:
        return view
    engine = get_ocr_engine()
    if isinstance(engine, NullOcrEngine) and _ocr_override is None:
        view.warnings.append("ocr_unavailable")
        if not text_is_sufficient(view.text):
            view.source = "empty"
        return view
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception:
        logger.exception("OCR could not reopen %s", filename or "resume")
        view.warnings.append("ocr_failed")
        return view
    original_digital = view.text
    ocr_text: dict[int, str] = {}
    confidences: list[float] = []
    try:
        if getattr(doc, "needs_pass", False):
            view.warnings.append("encrypted_pdf")
            view.source = "failed"
            return view
        for index in sparse[:8]:
            if index >= doc.page_count:
                break
            page = doc[index]
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            piece, confidence = _recognize_image(engine, pixmap.tobytes("png"))
            if confidence is not None:
                confidences.append(confidence)
            if confidence is not None and confidence < OCR_MIN_CONFIDENCE:
                view.warnings.append(f"ocr_low_confidence:{index + 1}")
                continue
            if text_is_sufficient(piece):
                ocr_text[index] = piece
                view.warnings.append(f"ocr_page:{index + 1}")
    except Exception:
        logger.exception("OCR failed for %s", filename or "resume")
        view.warnings.append("ocr_failed")
        return view
    finally:
        doc.close()
    if not ocr_text:
        if not text_is_sufficient(view.text):
            view.warnings.append("ocr_insufficient")
            if view.source != "failed":
                view.source = "empty"
        return view
    parts: list[str] = []
    blocks: list[TextBlock] = []
    for index in range(limit):
        if index in ocr_text:
            parts.append(ocr_text[index])
            for line_index, line in enumerate(ocr_text[index].splitlines()):
                cleaned = line.strip()
                if cleaned:
                    blocks.append(TextBlock(
                        text=cleaned,
                        page=index,
                        y0=float(line_index),
                        y1=float(line_index) + 1,
                        kind="ocr",
                    ))
            continue
        page_blocks = [block for block in view.blocks if block.page == index]
        blocks.extend(page_blocks)
        page_body = "\n".join(block.text for block in page_blocks if block.text.strip())
        if page_body:
            parts.append(page_body)
    view.text = clean_text("\n".join(parts))
    view.blocks = blocks
    view.ocr_pages = sorted(index + 1 for index in ocr_text)
    if confidences:
        view.ocr_confidence = round(sum(confidences) / len(confidences), 1)
    if not text_is_sufficient(original_digital):
        view.source = "ocr"
    return view


def _page_text(view: DocumentView, page_index: int) -> str:
    return "\n".join(block.text for block in view.blocks if block.page == page_index and block.text.strip())


def _extract_docx(filename: str, file_bytes: bytes) -> DocumentView:
    if not DOCX_AVAILABLE or Document is None:
        return DocumentView(source="failed", filename=filename, warnings=["docx_unavailable"])
    document = Document(io.BytesIO(file_bytes))
    blocks: list[TextBlock] = []
    y_cursor = 0.0

    def add_block(text: str, kind: str, x0: float = 0.0) -> None:
        nonlocal y_cursor
        cleaned = clean_text(text)
        if not cleaned:
            return
        blocks.append(TextBlock(text=cleaned, page=0, x0=x0, y0=y_cursor, x1=x0 + 240, y1=y_cursor + 12, kind=kind))
        y_cursor += 14

    for paragraph in _docx_header_footer_paragraphs(document):
        add_block(paragraph, "header")
    for paragraph in document.paragraphs:
        add_block(paragraph.text, "text")
    for table in document.tables:
        for row in table.rows:
            cells = [_cell_text(cell) for cell in row.cells]
            cells = [cell for cell in cells if cell]
            if not cells:
                continue
            add_block(" | ".join(cells), "table_cell")
    text = clean_text("\n".join(block.text for block in blocks))
    source = "docx" if text_is_sufficient(text) else ("empty" if not text else "docx")
    return DocumentView(text=text, blocks=blocks, source=source, filename=filename, page_count=1, columns=1)


def _docx_header_footer_paragraphs(document: Any) -> list[str]:
    texts: list[str] = []
    for section in document.sections:
        for container in (section.header, section.footer):
            try:
                for paragraph in container.paragraphs:
                    if paragraph.text.strip():
                        texts.append(paragraph.text)
            except Exception:
                continue
    return texts


def _cell_text(cell: Any) -> str:
    pieces = [paragraph.text.strip() for paragraph in cell.paragraphs if paragraph.text.strip()]
    if pieces:
        return clean_text(" ".join(pieces))
    if qn is None:
        return ""
    try:
        raw = cell._tc.xpath(".//w:t")
        return clean_text(" ".join(node.text or "" for node in raw))
    except Exception:
        return ""
